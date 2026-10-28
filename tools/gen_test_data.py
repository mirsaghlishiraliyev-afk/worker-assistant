"""Saxta test məlumatı: Test MMC, Sentyabr 2026. Uyğunlaşdırma və ƏDV hallarını ehtiva edir."""
import datetime as dt
import pathlib
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

OUT = pathlib.Path(__file__).resolve().parents[1] / "tests" / "data"
F = Font(name="Arial", size=10)
FB = Font(name="Arial", size=10, bold=True, color="FFFFFF")
HF = PatternFill("solid", fgColor="305496")


def D(d, m=9):
    return dt.date(2026, m, d)


kontr = [("1400000011", "Alfa Tikinti MMC"), ("1400000022", "Beta Logistika MMC"), ("1400000033", "Qamma Ticarət MMC"),
         ("1400000044", "Delta Servis MMC"), ("1400000055", "Eko Qida MMC"), ("1400000066", "Fokus Media MMC"),
         ("1400000077", "Günəş Enerji MMC"), ("1400000088", "Xəzər Texnika MMC"), ("1400000099", "İlkin Dizayn MMC"),
         ("1500000010", "Kür Avto MMC"), ("1500000021", "Lider Ofis MMC"), ("1500000032", "Mərcan Kimya MMC")]
name = dict(kontr)
TR = str.maketrans("əüöığşçƏÜÖİĞŞÇ", "euoigscEUOIGSC")
email = {v: f"{n.split()[0].translate(TR).lower()}@example.com" for v, n in kontr}

contracts = [("M-2026/01", "1400000011", D(10, 1), D(31, 12), 60000, "Tikinti materialları"),
             ("M-2026/02", "1400000022", D(1, 2), dt.date(2027, 1, 31), 40000, "Daşıma xidməti"),
             ("M-2025/17", "1400000033", dt.date(2025, 9, 1), D(31, 8), 30000, "Mal tədarükü"),
             ("M-2026/05", "1400000044", D(1, 3), dt.date(2027, 2, 28), 25000, "Texniki servis"),
             ("M-2026/07", "1400000055", D(1, 4), D(31, 12), 9000, "Qida tədarükü"),
             ("M-2026/09", "1400000066", D(1, 5), dt.date(2027, 4, 30), 20000, "Reklam xidməti"),
             ("M-2026/11", "1400000088", D(1, 6), dt.date(2027, 5, 31), 35000, "Avadanlıq"),
             ("M-2026/12", "1500000010", D(1, 7), D(31, 10), 15000, "Avto təmir"),
             ("A-2026/03", "1500000021", D(15, 1), D(31, 12), 12000, "Ofis ləvazimatı (alış)"),
             ("A-2026/04", "1500000032", D(1, 2), dt.date(2027, 1, 31), 18000, "Kimyəvi maddələr (alış)"),
             ("A-2026/06", "1400000022", D(1, 1), D(31, 12), 10000, "Nəqliyyat (alış)")]
prior = {"1400000055": 8200}

# --- Satış qaimələri: [seriya, nömrə, tarix, VÖEN, ƏDV-siz, ƏDV (None = formula), status]
plan = [("1400000011", [5000, 7200, 3100]), ("1400000022", [2500, 4000]), ("1400000033", [3500]),
        ("1400000044", [1800, 2600]), ("1400000055", [1500]), ("1400000066", [3000, 2200]),
        ("1400000077", [4100]), ("1400000088", [6500, 2900]), ("1500000010", [1200, 2400, 1200])]
sales, n, day = [], 1001, 2
for v, amts in plan:
    for a in amts:
        sales.append(["AA", n, D(day), v, a, None, "Təsdiqlənib"])
        n += 1
        day = day % 27 + 2
sales.append(list(sales[3]))                       # dublikat qaimə AA1004
sales[7][5] = round(sales[7][4] * 0.16, 2)         # AA1008: ƏDV 16%
S = {r[1]: r for r in sales[:-1]}

pplan = [("1500000021", [1200, 850, 640]), ("1500000032", [3400, 2100]), ("1400000022", [900, 1100]), ("1400000099", [2500])]
purch, m, day = [], 5001, 3
for v, amts in pplan:
    for a in amts:
        purch.append(["BB", m, D(day), v, a, None, "Təsdiqlənib"])
        m += 1
        day = day % 26 + 3
purch[2][6] = "Ləğv edilib"                         # BB5003

bank, dep = [], []


def vat(x):
    return round(x * 0.18, 2)


def pay(d, direction, v, amount, purpose):
    bank.append([d, direction, v, name[v], amount, purpose])


def vatpay(d, direction, v, amount, ref):
    dep.append([d, direction, v, name[v], amount, ref])


for no, r in S.items():
    sn, d, v, net = f"AA{no}", r[2], r[3], r[4]
    pd_ = min(d + dt.timedelta(days=2), D(30))
    if no == 1001:   # 2 qəpik yuvarlaşdırma fərqi
        pay(pd_, "Mədaxil", v, 4999.98, f"{sn} qaimə üzrə ödəniş"); vatpay(pd_, "Mədaxil", v, vat(net), sn)
    elif no == 1002:  # ödənilməyib
        pass
    elif no == 1003:  # qismən ödəniş
        pay(pd_, "Mədaxil", v, 2000, f"{sn} üzrə qismən ödəniş"); vatpay(pd_, "Mədaxil", v, vat(2000), sn)
    elif no == 1005:  # ƏDV depozitə gəlməyib
        pay(pd_, "Mədaxil", v, net, f"{sn} qaimə üzrə ödəniş")
    elif no == 1007:  # təyinatda nömrə yoxdur, məbləğ unikaldır
        pay(pd_, "Mədaxil", v, net, "Servis haqqı"); vatpay(pd_, "Mədaxil", v, vat(net), "")
    elif no == 1011:  # oktyabrda ödənilib
        pay(D(6, 10), "Mədaxil", v, net, f"{sn} qaimə üzrə ödəniş"); vatpay(D(6, 10), "Mədaxil", v, vat(net), sn)
    elif no in (1013, 1014):  # bir ödəniş iki qaiməni örtür
        if no == 1013:
            pay(D(24), "Mədaxil", v, 9400, "Avadanlıq üzrə ödəniş")
        vatpay(D(24), "Mədaxil", v, vat(net), sn)
    elif no == 1015:  # eyni məbləğli iki qaimə (1015, 1017), nömrəsiz təyinat
        pay(D(27), "Mədaxil", v, 1200, "Avto təmir ödənişi"); vatpay(D(27), "Mədaxil", v, vat(net), "")
    elif no == 1017:
        pass
    else:
        pay(pd_, "Mədaxil", v, net, f"{sn} qaimə üzrə ödəniş"); vatpay(pd_, "Mədaxil", v, vat(net), sn)

pay(D(18), "Mədaxil", "1400000077", 3800, "Xidmət haqqı")          # qaiməsiz mədaxil
beta = [b for b in bank if "AA1004" in b[5]][0]
bank.append(list(beta))                                               # bank sətri iki dəfə yüklənib

for i, r in enumerate(purch):
    sn, v, net = f"BB{r[1]}", r[3], r[4]
    d = min(r[2] + dt.timedelta(days=2), D(30))
    pay(d, "Məxaric", v, net, f"{sn} qaimə üzrə ödəniş")
    if i != 4:                                                        # BB5005: ƏDV depozitdən ödənilməyib
        vatpay(d, "Məxaric", v, vat(net), sn)

bank.sort(key=lambda x: x[0])
dep.sort(key=lambda x: x[0])


def sheet(ws, headers, widths):
    ws.append(headers)
    for c in ws[1]:
        c.font, c.fill = FB, HF
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"


def fin(ws):
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = F
            if isinstance(c.value, dt.date):
                c.number_format = "DD.MM.YYYY"
            elif isinstance(c.value, (int, float)) or (isinstance(c.value, str) and c.value.startswith("=")):
                c.number_format = "#,##0.00"


wb = Workbook()
o = wb.active
o.title = "Oxu"
info = [("SAXTA TEST MƏLUMATI", ""), ("Şirkət", "Test MMC (VÖEN 1300000001)"), ("Dövr", "Sentyabr 2026"),
        ("Məqsəd", "Uyğunlaşdırma mühərrikini yoxlamaq üçün. Bütün adlar, VÖEN-lər və məbləğlər uydurmadır."),
        ("Diqqət", "Sütun adları təxminidir. Real e-taxes və bank eksportlarının formatını mühasiblə yoxlayın."),
        ("ƏDV-nin ödənilməsi", "Alıcı ƏDV-siz məbləği bank hesabına, ƏDV-ni isə satıcının ƏDV depozit hesabına köçürür."),
        ("Vərəqlər", "E-qaimə_Satış, E-qaimə_Alış, Bank_çıxarışı, ƏDV_depozit, Müqavilələr, Kontragentlər"),
        ("Cavab açarı", "Ayrıca faylda: cavab_acari.xlsx"), ("", ""), ("ƏDV dərəcəsi", 0.18)]
for a, b in info:
    o.append([a, b])
o.column_dimensions["A"].width, o.column_dimensions["B"].width = 22, 95
for row in o.iter_rows():
    for c in row:
        c.font = F
        c.alignment = Alignment(wrap_text=True, vertical="top")
o["A1"].font = Font(name="Arial", size=12, bold=True, color="C00000")
o["B10"].number_format = "0%"
o["B10"].font = Font(name="Arial", size=10, color="0000FF")
o["C10"] = "Məlumat üçün. Mühərrikdə dərəcə uzlasma/config.py faylındadır."
o["C10"].font = F
o.column_dimensions["C"].width = 50


def inv_sheet(title, rows, party):
    ws = wb.create_sheet(title)
    sheet(ws, ["Seriya", "Nömrə", "Tarix", f"{party} VÖEN", f"{party} adı", "Məbləğ ƏDV-siz (AZN)", "ƏDV (AZN)", "Cəmi (AZN)", "Status"],
          [8, 9, 12, 14, 22, 18, 13, 14, 14])
    for i, r in enumerate(rows, 2):
        # Real e-taxes eksportu kimi: düsturlar yox, hazır dəyərlər
        v = r[5] if r[5] is not None else round(r[4] * 0.18, 2)
        ws.append([r[0], r[1], r[2], r[3], name[r[3]], r[4], v, round(r[4] + v, 2), r[6]])
    fin(ws)


inv_sheet("E-qaimə_Satış", sales, "Alıcı")
inv_sheet("E-qaimə_Alış", purch, "Satıcı")
ws = wb.create_sheet("Bank_çıxarışı")
sheet(ws, ["Tarix", "İstiqamət", "Kontragent VÖEN", "Kontragent adı", "Məbləğ (AZN)", "Təyinat"], [12, 11, 16, 22, 14, 30])
for r in bank:
    ws.append(r)
fin(ws)
ws = wb.create_sheet("ƏDV_depozit")
sheet(ws, ["Tarix", "İstiqamət", "Kontragent VÖEN", "Kontragent adı", "ƏDV məbləği (AZN)", "Qaimə"], [12, 11, 16, 22, 16, 12])
for r in dep:
    ws.append(r)
fin(ws)
ws = wb.create_sheet("Müqavilələr")
sheet(ws, ["Müqavilə №", "Kontragent VÖEN", "Kontragent adı", "Başlama", "Bitmə", "Məbləğ ƏDV-siz (AZN)", "Mövzu",
           "Yanvar–Avqust qaimələri (AZN)", "Növ"], [12, 16, 22, 12, 12, 18, 24, 20, 9])
for c in contracts:
    ws.append([c[0], c[1], name[c[1]], c[2], c[3], c[4], c[5], prior.get(c[1]), "Alış" if c[0].startswith("A-") else "Satış"])
fin(ws)
ws = wb.create_sheet("Kontragentlər")
sheet(ws, ["VÖEN", "Ad", "E-poçt"], [14, 24, 28])
for v, nm in kontr:
    ws.append([v, nm, email[v]])
fin(ws)
wb.save(f"{OUT}/test_meyar_data.xlsx")

key = [
    (1, "Satış qaiməsi ödənilməyib", "AA1002", "Alfa Tikinti: nə banka, nə ƏDV depozitə ödəniş gəlməyib", "Orta"),
    (2, "ƏDV depozitə gəlməyib", "AA1005", "Beta Logistika: ƏDV-siz hissə banka gəlib, ƏDV depozitə köçürülməyib", "Yüksək"),
    (3, "Qaiməsiz mədaxil", "Bank 18.09.2026", "Günəş Enerji: 3 800 AZN mədaxil var, uyğun açıq qaimə yoxdur", "Yüksək"),
    (4, "Dublikat qaimə", "AA1004", "Eyni seriya və nömrə satış siyahısında iki dəfə", "Orta"),
    (5, "Müqaviləsiz kontragent", "AA1012, BB5008", "Günəş Enerji və İlkin Dizayn üçün müqavilə yoxdur", "Orta"),
    (6, "Bitmiş müqavilə üzrə qaimə", "AA1006", "Qamma Ticarət: M-2025/17 31.08.2026-da bitib", "Orta"),
    (7, "Müqavilə məbləği aşılıb", "AA1009", "Eko Qida: limit 9 000, 8 200 + 1 500 = 9 700 AZN", "Orta"),
    (8, "ƏDV hesablanması səhvdir", "AA1008", "Qaimədə ƏDV 16%, depozitə 18% gəlib", "Yüksək"),
    (9, "Alış ƏDV-si depozitdən ödənilməyib", "BB5005", "Mərcan Kimya: ƏDV hissəsi ödənilməyib", "Yüksək"),
    (10, "Ləğv edilmiş qaimə üzrə ödəniş", "BB5003", "Qaimə ləğv edilib, ödəniş aparılıb", "Orta"),
    (11, "Dövr fərqi", "AA1011", "Ödəniş oktyabrda (06.10.2026) gəlib", "Aşağı"),
    (12, "Müddəti bitmək üzrə müqavilə", "M-2026/12", "Kür Avto: 31.10.2026-da bitir", "Məlumat"),
    (13, "Yuvarlaşdırma fərqi (uyğun sayılmalı)", "AA1001", "4 999,98 ödənilib, qaimə 5 000: 2 qəpik fərq, dözüm daxilində", "Yoxdur"),
    (14, "Qismən ödəniş", "AA1003", "3 100-dən 2 000 ödənilib, 1 100 qalıq", "Orta"),
    (15, "Nömrəsiz təyinat, unikal məbləğ", "AA1007", "'Servis haqqı': VÖEN və məbləğə görə uyğunlaşdırılmalı", "Yoxdur"),
    (16, "Bir ödəniş iki qaimə", "AA1013 + AA1014", "9 400 = 6 500 + 2 900", "Yoxdur"),
    (17, "Qeyri-müəyyən uyğunluq", "AA1015 / AA1017", "Kür Avto: 1 200 ödəniş, eyni məbləğli iki qaimə, mühasib seçməlidir", "Orta"),
    (18, "Bank sətri təkrar", "Bank, AA1004", "Beta Logistika ödənişi çıxarışda iki dəfə", "Orta"),
]
kb = Workbook()
k = kb.active
k.title = "Cavab_açarı"
sheet(k, ["#", "Hal", "İstinad", "İzah", "Risk"], [5, 34, 20, 80, 10])
for r in key:
    k.append(list(r))
for row in k.iter_rows(min_row=2):
    for c in row:
        c.font = F
        c.alignment = Alignment(wrap_text=True, vertical="top")
kb.save(f"{OUT}/cavab_acari.xlsx")
print(len(sales), len(purch), len(bank), len(dep))
