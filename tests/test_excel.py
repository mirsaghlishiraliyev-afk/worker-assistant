"""İstənilən Excel faylının oxunması və təhlili (SQLite keşi, başlıq sətri, mətn kimi yazılmış rəqəmlər, problemlər)."""
import pathlib
import sys
import tempfile

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parents[1]))
from assistant.modullar import excel as EX  # noqa: E402

FAYL = pathlib.Path(__file__).parent / "data" / "ixtiyari_excel.xlsx"
KES = tempfile.mkdtemp()


class S:
    def __init__(self):
        self.excel = EX.ac(FAYL, KES)


def test_basliq_ve_tipler():
    v = S().excel["vereqler"]["Satışlar 2026"]
    assert v["sutunlar"][:2] == ["Tarix", "Müştəri"]           # başlıq 4-cü sətirdədir
    assert v["basliq_setri"] == 4
    assert v["tipler"]["Məbləğ (AZN)"] == "rəqəm"               # "1 234,56" yazılışları da rəqəmdir
    assert v["tipler"]["Tarix"] == "tarix"


def test_kes_ikinci_defe_cevirmir():
    a = EX.ac(FAYL, KES)
    b = EX.ac(FAYL, KES)
    assert a["db"] == b["db"]


def test_problemler():
    m = "".join(EX.problemler(S(), vereq=v)["metn"] for v in ["Satışlar 2026", "İşçilər"])
    for soz in ["Təkrar sətirlər", "Qeyri-adi dəyərlər", "Mənfi dəyərlər", "Rəqəm sütununda mətn", "Şübhəli tarixlər", "68"]:
        assert soz in m, soz


def test_lokal_sorgular():
    s = S()
    assert EX.qayda(s, "Məbləğ cəmi müştəri üzrə")[1]["qrup_sutunu"] == "Müştəri"
    assert EX.qayda(s, "Ən kiçik 3 maaş")[1]["vereq"] == "İşçilər"
    assert EX.qayda(s, "'Gəncə' axtar") == ("axtar", {"sorgu": "Gəncə"})


def test_qruplama_neticesi():
    n = EX.qrupla(S(), "Məbləğ", "Müştəri", vereq="Satışlar 2026")
    assert n["ui"]["setirler"][0] == ["Alfa Tikinti", "109 574,19"]


def test_obyekt_suallari():
    s = S()
    plan = EX.qayda(s, "Aynurun maaşı nə qədərdir?")
    assert plan[0] == "obyekt_gostericisi"
    netice = EX.obyekt_gostericisi(s, **plan[1])
    assert "Maaş" in netice["metn"] and "1 500" in netice["metn"]
    netice = EX.obyekt_gostericisi(s, "Alfa Tikinti", "satışların məbləği")
    assert "109 574,19" in netice["metn"]


def test_movzu_ixtiyari_vereqe_dusmur():
    s = S()
    plan = EX.qayda(s, "işçiləri göstər")
    netice = EX.movzu_goster(s, **plan[1]) if plan[0] == "movzu_goster" else EX.setirler(s, **plan[1])
    assert "«İşçilər»" in netice["metn"]
    assert EX.movzu_goster(s, "avansları göstər")["ui"] is None   # belə vərəq yoxdur: ixtiyari vərəq göstərilmir


def _jurnal_fayli():
    """Kiçik 1C-yə bənzər fayl: müxabirləşmə jurnalı + hesablar planı."""
    import openpyxl
    wb = openpyxl.Workbook()
    j = wb.active
    j.title = "Müxabirləşmə jurnalı"
    j.append(["Dövr", "Hesab Dt", "Hesab Ct", "Məbləğ", "Mündəricat"])
    import datetime as dt
    for row in [("2025-03-01", 721, 531, 1000.0, "icarə"), ("2025-04-01", 721, 223, 500.5, "kommunal"),
                ("2026-01-10", 731, 223, 200.0, "məzənnə"), ("2026-02-01", 223, 601, 5000.0, "gəlir"),
                ("2026-02-02", 531, 223, 300.0, "ödəniş"), ("2026-03-01", 721, 536.05, 80.0, "ekspert rəyi"),
                ("2026-03-05", 536.05, 607.05, 80.0, "ekspert")]:
        j.append([dt.datetime.fromisoformat(row[0]), *row[1:]])
    p = wb.create_sheet("Mühasibat uçotunun hesablar planı")
    p.append(["Kod", "Ad"])
    for row in [(721, "İnzibati xərclər"), (731, "Məzənnə fərqi"), (601, "Satış"), (531, "Əməyin ödənilməsi"),
                (536.05, "Ekspertlərlə hesablaşmalar"), (607.05, "Müstəqil ekspertlərlə hesablaşmalar")]:
        p.append(list(row))
    a = wb.create_sheet("Avans hesabatı")
    a.append(["Tarix", "Əməkdaş", "Sənədin məbləği"])
    a.append(["2026-01-05", "X", 40.0])
    a.append(["2026-01-06", "Y", 60.0])
    yol = pathlib.Path(KES) / "jurnal.xlsx"
    wb.save(yol)
    s = S.__new__(S)
    s.excel, s.son_niyyet = EX.ac(yol, KES), None
    return s


def test_umumi_xerc_jurnaldan():
    s = _jurnal_fayli()
    plan = EX.qayda(s, "Şirkətin ümumi xərci nə qədərdir?")
    assert plan[0] == "cem_hesabla"
    n = EX.cem_hesabla(s, **plan[1])
    assert "1 780,50" in n["metn"] and "Debet 7xx" in n["metn"]          # 1000 + 500,50 + 200 + 80; 531 daxil deyil
    assert n["ui"]["setirler"][0][:2] == ["531", "Əməyin ödənilməsi"]  # xərc növü = qarşı hesab
    assert "**280**" in EX.cem_hesabla(s, "2026-cı ildə ümumi xərc")["metn"]
    assert "5 000" in EX.cem_hesabla(s, "ümumi gəlir")["metn"]


def test_movzu_cemi_ve_davam():
    s = _jurnal_fayli()
    assert "100" in EX.cem_hesabla(s, "ümumi avans məbləği")["metn"]
    s2 = S()
    s2.son_niyyet = None
    n = EX.cem_hesabla(s2, "ümumi tutulma cəmi")                           # belə vərəq yoxdur: soruşur
    assert "hansı vərəqi" in n["metn"] and s2.son_niyyet
    plan = EX.qayda(s2, "Satışlar 2026")                                   # cavab: vərəqin adı -> cəmi hesablanır
    assert plan[0] == "cem_hesabla" and plan[1]["vereq"] == "Satışlar 2026"
    assert "**" in EX.cem_hesabla(s2, **plan[1])["metn"]


def test_xerclerin_icinde_ekspertler_ve_ayar():
    s = _jurnal_fayli()
    EX.AYAR_FAYLI = pathlib.Path(KES) / "ayar.json"
    plan = EX.qayda(s, "Xərclərin içində ekspertlər nə qədərdir?")
    n = EX.cem_hesabla(s, **plan[1])
    assert "**80**" in n["metn"] and "Ekspertlərlə hesablaşmalar" in n["ui"]["setirler"][0][1]
    assert "607.05" in n["metn"] and "təkrar sayıla bilər" in n["metn"]   # 6-cı bölmədəki ekspert hesabı barədə xəbərdarlıq
    plan = EX.qayda(s, "607.05 hesabını xərclərə daxil et")
    assert plan == ("hesab_ayarla", {"hesab": "607.05", "bolme": "xerc", "emeliyyat": "daxil"})
    EX.hesab_ayarla(s, **plan[1])
    assert "1 860,50" in EX.cem_hesabla(s, "ümumi xərc")["metn"]           # 607.05-in kredit dövriyyəsi xərcə əlavə olunur
    EX.hesab_ayarla(s, "607.05", emeliyyat="cixar")
    assert "1 780,50" in EX.cem_hesabla(s, "ümumi xərc")["metn"]


def test_gelir_hesablasma_hesabi_daxil_deyil():
    s = _jurnal_fayli()
    EX.AYAR_FAYLI = pathlib.Path(KES) / "ayar_gelir.json"
    n = EX.cem_hesabla(s, "Şirkətin ümumi gəliri")
    assert "**5 000**" in n["metn"]                                     # 607.05 «...hesablaşmalar» gəlir deyil
    assert "607.05" in n["metn"] and "gəlir deyil" in n["metn"]
    assert [r[0] for r in n["ui"]["setirler"]] == ["601", "Cəmi"]       # gəlir növü = gəlir hesabı, qarşı hesab yox
    EX.hesab_ayarla(s, "607.05", "gelir", "daxil")
    assert "**5 080**" in EX.cem_hesabla(s, "ümumi gəlir")["metn"]
    assert EX.qayda(s, "607.05 gəlir deyil")[1] == {"hesab": "607.05", "bolme": "gelir", "emeliyyat": "cixar"}
    EX.hesab_ayarla(s, **EX.qayda(s, "607.05 gəlir deyil")[1])
    assert "**5 000**" in EX.cem_hesabla(s, "ümumi gəlir")["metn"]


def test_hesab_xercdir_emri_adla():
    s = _jurnal_fayli()
    EX.AYAR_FAYLI = pathlib.Path(KES) / "ayar_603.json"
    plan = EX.qayda(s, "Satış xərcdir")                                    # hesab adı ilə, kodsuz
    assert plan == ("hesab_ayarla", {"hesab": "601", "bolme": "xerc", "emeliyyat": "daxil"})
    n = EX.hesab_ayarla(s, **plan[1])
    assert "fərqlənəcək" in n["metn"]                                      # 601 1C-də gəlir kimi bağlanır: xəbərdarlıq
    n = EX.cem_hesabla(s, "ümumi xərc")
    assert "**6 780,50**" in n["metn"]                                       # 1 780,50 + 5 000
    assert ["601", "Satış", "5 000", "1"] in n["ui"]["setirler"]
    assert "Satış" in EX.cem_hesabla(s, "ümumi gəlir")["metn"]             # gəlirdən çıxarılıb
    plan = EX.qayda(s, "601 hesabını xərcdən çıxar")
    assert plan[1] == {"hesab": "601", "bolme": "xerc", "emeliyyat": "cixar"}
    EX.hesab_ayarla(s, **plan[1])
    assert "**5 000**" in EX.cem_hesabla(s, "ümumi gəlir")["metn"]         # gəlirə qayıtdı
    assert EX.qayda(s, "Təkrar sətirləri tap")[0] == "problemler"
