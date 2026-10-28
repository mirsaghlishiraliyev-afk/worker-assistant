"""Bank daxilolmaları: çıxarışın tanınması, təsnifat, öyrənmə, xəbərdarlıqlar, 1C faylı."""
import pathlib
import sys
import tempfile

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parents[1]))
from assistant.modullar import bank as B  # noqa: E402
from assistant.modullar import excel as EX  # noqa: E402

FAYL = pathlib.Path(__file__).parent / "data" / "bank_cixarisi_test.xlsx"
QOV = pathlib.Path(tempfile.mkdtemp())


class S:
    def __init__(self, ad="q.json"):
        B.QAYDA_FAYLI = QOV / ad
        B.HESABAT_QOVLUGU = QOV
        self.bank = B.oxu(FAYL)


def _kat(s, setir):
    return next(r for r in s.bank["setirler"] if r["setir"] == setir)["kateqoriya"]


def test_tanima_ve_oxuma():
    s = S()
    assert s.bank["basliq_setri"] == 4
    assert len(s.bank["setirler"]) == 12                      # debet və «Cəmi» sətri atılır
    assert round(sum(r["mebleg"] for r in s.bank["setirler"]), 2) == 27608.45  # «3 400,00» mətn kimi yazılıb
    assert B.oxu(pathlib.Path(__file__).parent / "data" / "ixtiyari_excel.xlsx") is None


def test_tesnifat():
    s = S()
    assert [_kat(s, i) for i in (5, 6, 7)] == ["sigorta_haqqi"] * 3
    assert _kat(s, 9) == "subroqasiya"
    assert _kat(s, 10) == "emek_haqqi" and _kat(s, 11) == "ezamiyye"
    assert _kat(s, 12) == "yoxlanmali"                         # xəzinədarlıq, amma təyinat aydın deyil
    assert _kat(s, 13) == "telim"
    assert _kat(s, 8) == "yoxlanmali"                          # «Hesaba mədaxil»
    r = next(r for r in s.bank["setirler"] if r["setir"] == 5)
    assert r["muqavile"] == "HS-2026000123-QS"


def test_xeberdarliqlar():
    s = S()
    nov = {(a, b) for a, b, _ in s.bank["xeberdarliqlar"]}
    assert (15, "Başqa ay") in nov
    assert (16, "Təkrar ola bilər") in nov


def test_oyrenme_ve_madde():
    s = S("oyren.json")
    plan = B.qayda(s, "8-ci sətir sığorta haqqıdır")
    assert plan == ("bank_teyin", {"setir": 8, "kateqoriya": "sigorta_haqqi"})
    n = B.bank_teyin(s, **plan[1])
    assert "daha 1" in n["metn"]                                # eyni VÖEN-li 17-ci sətir də keçdi
    assert _kat(s, 17) == "sigorta_haqqi"
    assert _kat(S("oyren.json"), 8) == "sigorta_haqqi"         # yeni çıxarışda da tanınır
    plan = B.qayda(s, 'Subroqasiya maddəsi "Subroqasiya gəlirləri" olsun')
    assert plan == ("bank_madde", {"kateqoriya": "subroqasiya", "madde": "Subroqasiya gəlirləri"})
    B.bank_madde(s, **plan[1])
    assert B._madde("subroqasiya") == "Subroqasiya gəlirləri"


def test_1c_fayli():
    import openpyxl
    s = S("hesabat.json")
    n = B.bank_hesabat(s)
    wb = openpyxl.load_workbook(QOV / n["ui"]["fayllar"][1]["ad"])   # yoxlama faylı
    ws = wb["1C - Sığorta haqları"]
    assert [c.value for c in ws[1]][:len(B.SUTUNLAR_1C)] == B.SUTUNLAR_1C
    assert ws.max_row - 1 == 2                                  # yalnız HS-/BS- müqavilə nömrəsi olanlar (5 və 7-ci sətir)
    assert wb["Sığorta haqları (müqaviləsiz)"].max_row - 1 == 2
    assert any(c.value == 16 for c in wb["Yoxlanmalı"]["A"])
    assert "Yoxlanmalı" in wb.sheetnames and "Xəbərdarlıqlar" in wb.sheetnames


def test_lokal_sorgular():
    s = S()
    assert B.qayda(s, "Yoxlanmalıları göstər") == ("bank_siyahi", {"kateqoriya": "yoxlanmali"})
    assert B.qayda(s, "Subroqasiya daxilolmaları") == ("bank_siyahi", {"kateqoriya": "subroqasiya"})
    assert B.qayda(s, "1C üçün Excel hazırla") == ("bank_hesabat", {})
    assert B.qayda(s, "Xülasə") == ("bank_xulase", {})
    assert EX is not None


def test_hazir_sorgular():
    from assistant import teklif as TK
    s = S("teklif.json")
    s.excel, s.m, s.son_niyyet = {"ad": "x", "vereqler": {}}, None, None
    S.hazir = property(lambda self: self.m is not None)
    t = TK.hazirla(s)
    assert t[0] == "Sualları ver (4)" and "1C üçün Excel hazırla" in t      # 3 tanınmayan + 1 «Orta» (subroqasiya)
    assert all(B.qayda(s, x) for x in t)                       # hər təklif başa düşülür


KAPITAL = pathlib.Path(__file__).parent / "data" / "bank_kapital_test.xlsx"


def test_kapital_formati():
    B.QAYDA_FAYLI, B.HESABAT_QOVLUGU = QOV / "kapital.json", QOV
    b = B.oxu(KAPITAL)
    assert b and b["meta"]["hesab"] == "AZ00TEST0000000000000000000"
    ss = {r["setir"]: r for r in b["setirler"]}
    assert len(ss) == 9                                            # mənfi (məxaric) sətir atılır
    k = {i: r["kateqoriya"] for i, r in ss.items()}
    assert k[8] == "telim" and k[10] == "telim"                    # «AQRAR SIĞORTA FONDU» təyinatda sığorta sayılmır
    assert k[9] == "sigorta_haqqi" and k[11] == "sigorta_haqqi"
    assert k[12] == "ezamiyye"                                     # xəzinədarlıq hesabı (CTRE)
    assert k[14] == "qaytarilma" and k[16] == "faiz" and k[17] == "depozit" and k[15] == "yoxlanmali"
    assert ss[9]["odeyici"] == "NUMUNOVA LALE KAMIL QIZI" and ss[9]["fin"] == "1AB2CD3" and ss[9]["teyinat"] == "SIĞORTA"
    assert ss[8]["odeyici"] == "TESTOV AYDIN VAQIF OGLU" and ss[8]["voen"] == ""   # ödəyici bank deyil, şəxsdir
    assert ss[11]["odeyici"] == "Testov Elnur" and ss[11]["muqavile"] == "HS-2026000001-QS"
    assert ss[11]["t_fin_voen"] == "1234567890" and ss[11]["odeyici_hesab"] == "AZ00RBTA00000000000000000099"


def test_1c_inteqrasiya_sablonu():
    import openpyxl
    B.QAYDA_FAYLI, B.HESABAT_QOVLUGU = QOV / "kapital2.json", QOV

    class K:
        bank = B.oxu(KAPITAL)
    n = B.bank_hesabat(K())
    f1c = n["ui"]["fayllar"][0]["ad"]
    wb = openpyxl.load_workbook(QOV / f1c)
    assert wb.sheetnames == ["Sheet1"]
    rows = list(wb["Sheet1"].iter_rows(values_only=True))
    assert list(rows[0]) == ["TARİX", "Sığortalı Fin/Voen", "Sığorta müqaviləsinin nömrəsi", "HESAB NÖMRƏSİ", "MƏBLƏĞ", "Pul vəsaitinin maddəsi"]
    assert len(rows) == 2                                       # yalnız HS- nömrəli sətir
    assert rows[1][1:5] == ("1234567890", "HS-2026000001-QS", "AZ00RBTA00000000000000000099", 2700)
    assert rows[1][5] == "Sığorta haqları"


def test_muqavilesizler_istege_gore():
    import openpyxl
    B.QAYDA_FAYLI, B.HESABAT_QOVLUGU = QOV / "kapital3.json", QOV

    class K:
        bank = B.oxu(KAPITAL)
    plan = B.qayda(K(), "Müqaviləsizləri də şablona sal")
    assert plan == ("bank_hesabat", {"muqavilesiz": True})
    n = B.bank_hesabat(K(), **plan[1])
    rows = list(openpyxl.load_workbook(QOV / n["ui"]["fayllar"][0]["ad"])["Sheet1"].iter_rows(values_only=True))
    assert len(rows) == 3                                       # HS- nömrəli + müqaviləsiz sığorta haqqı
    assert {r[2] for r in rows[1:]} == {"HS-2026000001-QS", None}
    assert B.qayda(K(), "Müqaviləsizləri salma")[1] == {"muqavilesiz": False}


def test_muhasibe_sual_ve_cavab():
    s = S("sual.json")
    plan = B.qayda(s, "Sualları ver")
    assert plan == ("bank_sual", {})
    n = B.bank_sual(s)
    assert n["metn"].startswith("Sual 1/4: 8-ci sətir") and n["ui"]["tip"] == "sual"
    assert any(v["mesaj"] == "8-ci sətir — Sığorta haqları olsun" for v in n["ui"]["variantlar"])
    # düymə ilə cavab: təyin edir və növbəti sualı verir
    n = B.bank_teyin(s, **B.qayda(s, "8-ci sətir — Sığorta haqları olsun")[1])
    assert "Oldu" in n["metn"] and "Sual" in n["metn"] and n["ui"]["setir"] != 8
    # sərbəst cavab aktiv suala aiddir
    aktiv = s.bank["aktiv_sual"]
    assert B.qayda(s, "bilmirəm, sonra baxaram") == ("bank_kec", {"setir": aktiv})
    n = B.bank_kec(s, aktiv)
    aktiv = s.bank["aktiv_sual"]
    assert B.qayda(s, "bu digər daxilolmadır") == ("bank_teyin", {"setir": aktiv, "kateqoriya": "diger"})
    # «Orta» təxmin üçün «bəli» təxmini təsdiqləyir
    while s.bank["aktiv_sual"] is not None and next(x for x in s.bank["setirler"] if x["setir"] == s.bank["aktiv_sual"])["kateqoriya"] == "yoxlanmali":
        B.bank_teyin(s, s.bank["aktiv_sual"], "diger")
    if s.bank["aktiv_sual"] is not None:
        r = next(x for x in s.bank["setirler"] if x["setir"] == s.bank["aktiv_sual"])
        assert B.qayda(s, "Bəli") == ("bank_teyin", {"setir": r["setir"], "kateqoriya": r["kateqoriya"]})


def test_mushasibin_sablonunu_doldurur():
    import datetime as dt
    import openpyxl
    B.QAYDA_FAYLI, B.HESABAT_QOVLUGU = QOV / "sablon.json", QOV
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(B.SUTUNLAR_1C)
    ws.append([dt.datetime(2026, 9, 30), "1111111111", "HS-2026000000-QS", "AZ00", 10.0, "Sığorta haqları"])
    yol = QOV / "Sığorta haqqı inteqrasiya faylı.xlsx"
    wb.save(yol)
    assert B.sablon_mi(yol)[2]["MƏBLƏĞ"] == 5
    assert B.sablon_mi(KAPITAL) is None and B.sablon_mi(FAYL) is None

    class K:
        bank = B.oxu(KAPITAL)
        sablon = {"yol": yol, "ad": yol.name}
    n = B.bank_hesabat(K())
    assert n["metn"].startswith("«Sığorta haqqı inteqrasiya faylı.xlsx» şablonunuz dolduruldu: 1 yeni sətir")
    rows = list(openpyxl.load_workbook(QOV / n["ui"]["fayllar"][0]["ad"])["Sheet1"].iter_rows(values_only=True))
    assert len(rows) == 3 and rows[1][2] == "HS-2026000000-QS" and rows[2][2] == "HS-2026000001-QS"   # köhnə sətir qalır
