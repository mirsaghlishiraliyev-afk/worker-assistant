"""Mənbə faylları oxuyub standart cədvəllərə çevirir və dublikatları işarələyir."""
import hashlib

import pandas as pd

from . import config as C


class FaylXetasi(ValueError):
    """İstifadəçiyə göstərilə bilən, başa düşülən xəta."""


HESABAT_VEREQLERI = {"Xülasə", "İstisnalar", "Bağlantılar"}


def yoxla(fayl):
    """Faylın giriş faylı olduğunu və lazım olan vərəqlərin, sütunların olduğunu yoxlayır."""
    from openpyxl import load_workbook
    try:
        wb = load_workbook(fayl, read_only=True)
    except Exception:
        raise FaylXetasi("Fayl Excel kimi açılmadı. Faylın zədələnmədiyini və .xlsx formatında olduğunu yoxlayın.")
    vereqler = set(wb.sheetnames)
    if HESABAT_VEREQLERI <= vereqler:
        raise FaylXetasi("Bu, Worker's Assistant-ın hazırladığı hesabatdır, giriş faylı deyil. "
                         "Qaimələr, bank çıxarışı və ƏDV depozit vərəqləri olan faylı yükləyin.")
    lazim = [m["vereq"] for m in C.MENBELER.values()]
    catismir = [v for v in lazim if v not in vereqler]
    if catismir:
        raise FaylXetasi("Faylda bu vərəqlər tapılmadı: " + ", ".join(catismir) +
                         ". Lazım olan vərəqlər: " + ", ".join(lazim) + ".")
    for menbe, m in C.MENBELER.items():
        ws = wb[m["vereq"]]
        basliq = {str(c.value).strip() for c in next(ws.iter_rows(min_row=1, max_row=1)) if c.value is not None}
        yox = [s for s in m["sutunlar"].values() if s not in basliq]
        if yox:
            raise FaylXetasi(f"«{m['vereq']}» vərəqində bu sütunlar yoxdur: " + ", ".join(yox) + ".")
    wb.close()


QEYRI_METN = {"tarix", "netto", "edv", "mebleg", "baslama", "bitme", "evvelki"}


def _oxu(fayl, menbe):
    m = C.MENBELER[menbe]
    df = pd.read_excel(fayl, sheet_name=m["vereq"])
    ters = {v: k for k, v in m["sutunlar"].items()}
    df = df.rename(columns=ters)[list(m["sutunlar"].keys())]
    for c in df.columns:
        if c not in QEYRI_METN:
            df[c] = df[c].where(df[c].notna(), "").astype(str).str.strip().str.replace(r"\.0$", "", regex=True)
    return df


def _reqem(s):
    if s.dtype == object:
        s = s.str.replace(" ", "", regex=False).str.replace(",", ".", regex=False).replace("", "0")
    return pd.to_numeric(s, errors="coerce").fillna(0.0).round(2)


def _tarix(s):
    if pd.api.types.is_datetime64_any_dtype(s):
        return s.dt.date
    return pd.to_datetime(s, errors="coerce", format="mixed", dayfirst=True).dt.date


def qaimeler(fayl):
    hisseler = []
    for menbe, nov in (("satis", "Satış"), ("alis", "Alış")):
        df = _oxu(fayl, menbe)
        df["id"] = df["seriya"].str.upper() + df["nomre"]
        df["nov"] = nov
        df["tarix"] = _tarix(df["tarix"])
        df["netto"] = _reqem(df["netto"])
        df["edv"] = _reqem(df["edv"])
        hisseler.append(df)
    q = pd.concat(hisseler, ignore_index=True)
    q["dublikat"] = q.duplicated(subset=["nov", "id"], keep="first")
    q["legv"] = q["status"].isin(C.LEGV_STATUSLARI)
    return q


def odenisler(fayl):
    hisseler = []
    for axin in ("bank", "depozit"):
        df = _oxu(fayl, axin)
        df["axin"] = axin
        df["tarix"] = _tarix(df["tarix"])
        df["mebleg"] = _reqem(df["mebleg"])
        hisseler.append(df)
    o = pd.concat(hisseler, ignore_index=True)
    o["nov"] = o["istiqamet"].map(C.ISTIQAMET)
    sahələr = ["axin", "tarix", "istiqamet", "voen", "mebleg", "metn"]
    o["hash"] = ["|".join(map(str, r)) for r in o[sahələr].itertuples(index=False)]
    o["hash"] = o["hash"].map(lambda s: hashlib.sha1(s.encode()).hexdigest()[:12])
    o["dublikat"] = o.duplicated(subset=["hash"], keep="first")
    o["id"] = [f"{'B' if a == 'bank' else 'D'}{i + 1:03d}" for i, a in enumerate(o["axin"])]
    return o


def muqavileler(fayl):
    df = _oxu(fayl, "muqavile")
    df["baslama"] = _tarix(df["baslama"])
    df["bitme"] = _tarix(df["bitme"])
    df["mebleg"] = _reqem(df["mebleg"])
    df["evvelki"] = _reqem(df["evvelki"])
    return df


def kontragentler(fayl):
    return _oxu(fayl, "kontragent")
