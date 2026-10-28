"""Nəticələri mühasib üçün Excel hesabatına yazır."""
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from . import config as C
from .config import az

F = Font(name="Arial", size=10)
FB = Font(name="Arial", size=10, bold=True, color="FFFFFF")
HF = PatternFill("solid", fgColor="305496")
RISK_RENG = {"Yüksək": "F8CBAD", "Orta": "FFE699", "Aşağı": "E2EFDA", "Məlumat": "DDEBF7"}
RISK_SIRA = {"Yüksək": 0, "Orta": 1, "Aşağı": 2, "Məlumat": 3}


def _vereq(wb, ad, basliqlar, enler, setirler, reqem_sutunlari=(), tarix_sutunlari=()):
    ws = wb.create_sheet(ad)
    ws.append(basliqlar)
    for c in ws[1]:
        c.font, c.fill = FB, HF
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for i, w in enumerate(enler, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for s in setirler:
        ws.append(s)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = F
            c.alignment = Alignment(vertical="top", wrap_text=True)
            if c.column in reqem_sutunlari:
                c.number_format = "#,##0.00"
            if c.column in tarix_sutunlari:
                c.number_format = "DD.MM.YYYY"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    return ws


def borclar(m, dovr_sonu):
    """Dövrün sonuna olan qalıq: dövrdən sonrakı ödənişlər nəzərə alınmır."""
    sonra = {}
    for x in m.baglantilar:
        if x["odenis_tarixi"] > dovr_sonu:
            k = (x["axin"], x["qaime_id"])
            sonra[k] = sonra.get(k, 0) + x["baglanan_mebleg"]
    setirler = []
    for _, r in m.q[~m.q["legv"]].iterrows():
        qn = m.qalan[("bank", r["id"])] + sonra.get(("bank", r["id"]), 0)
        qe = m.qalan[("depozit", r["id"])] + sonra.get(("depozit", r["id"]), 0)
        qal = round(max(qn, 0) + max(qe, 0), 2)
        if qal > C.DOZUM:
            cem = round(r["netto"] + r["edv"], 2)
            setirler.append({"nov": r["nov"], "voen": r["voen"], "ad": r["ad"], "qaime": r["id"], "tarix": r["tarix"],
                             "cem": cem, "odenilib": round(cem - qal, 2), "qaliq": qal})
    return pd.DataFrame(setirler)


def yaz(fayl, m, istisnalar, odenisler, kontragentler, dovr_sonu, tesdiq=None):
    """tesdiq: (odenis_id, qaime_id) -> "Bəli" / "Xeyr" / "" qaytaran funksiya (istəyə bağlı)."""
    wb = Workbook()
    wb.remove(wb.active)
    b = pd.DataFrame(borclar(m, dovr_sonu))
    o = m.o
    bagli = {x["odenis_id"] for x in m.baglantilar}
    avto = len(bagli)
    cem_od = len(o)

    # Xülasə
    ws = _vereq(wb, "Xülasə", ["Göstərici", "Dəyər"], [48, 22], [
        ["Şirkət", C.SIRKET_ADI],
        ["Dövrün sonu", dovr_sonu],
        ["Ödəniş sətirləri (bank + depozit, təkrarsız)", cem_od],
        ["Avtomatik bağlanan ödənişlər", avto],
        ["Avtomatik bağlanma faizi", avto / cem_od if cem_od else 0],
        ["İstisnalar (mühasib baxmalıdır)", len(istisnalar)],
        ["  o cümlədən yüksək riskli", sum(1 for i in istisnalar if i["risk"] == "Yüksək")],
        ["Debitor borc, AZN (bizə borclular)", float(b[b["nov"] == "Satış"]["qaliq"].sum()) if len(b) else 0],
        ["Kreditor borc, AZN (bizim borclarımız)", float(b[b["nov"] == "Alış"]["qaliq"].sum()) if len(b) else 0],
    ])
    ws["B3"].number_format = "DD.MM.YYYY"
    ws["B6"].number_format = "0%"
    ws["B9"].number_format = ws["B10"].number_format = "#,##0.00"
    ws.append([])
    ws.append(["Qeyd: sistem yalnız təklif edir. 'Bağlantılar' vərəqində hər sətri təsdiqləyin və ya rədd edin."])
    ws.cell(ws.max_row, 1).font = Font(name="Arial", size=10, italic=True)

    # İstisnalar
    ist = sorted(istisnalar, key=lambda i: (RISK_SIRA.get(i["risk"], 9), i["nov"]))
    ad = dict(zip(kontragentler["voen"], kontragentler["ad"]))
    ws = _vereq(wb, "İstisnalar", ["Risk", "Növ", "İstinad", "Kontragent", "İzah", "Mühasibin qeydi"],
                [10, 30, 20, 22, 70, 30],
                [[i["risk"], i["nov"], i["istinad"], ad.get(i["voen"], i["voen"]), i["izah"], ""] for i in ist])
    for row in ws.iter_rows(min_row=2):
        rg = RISK_RENG.get(row[0].value)
        if rg:
            row[0].fill = PatternFill("solid", fgColor=rg)

    # Bağlantılar
    ws = _vereq(wb, "Bağlantılar",
                ["Ödəniş", "Axın", "Tarix", "Kontragent", "Ödəniş məbləği", "Təyinat", "Qaimə", "Bağlanan məbləğ",
                 "Qayda", "Etibar", "Təsdiq (Bəli/Xeyr)"],
                [8, 9, 11, 22, 14, 26, 10, 14, 26, 9, 14],
                [[x["odenis_id"], "Bank" if x["axin"] == "bank" else "ƏDV depozit", x["odenis_tarixi"], x["kontragent"],
                  x["odenis_meblegi"], x["metn"], x["qaime_id"], x["baglanan_mebleg"], x["qayda"], x["etibar"],
                  tesdiq(x["odenis_id"], x["qaime_id"]) if tesdiq else ""]
                 for x in m.baglantilar], reqem_sutunlari=(5, 8), tarix_sutunlari=(3,))
    dv = DataValidation(type="list", formula1='"Bəli,Xeyr"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"K2:K{max(ws.max_row, 2)}")
    for row in ws.iter_rows(min_row=2):
        if row[9].value == "Orta":
            row[9].fill = PatternFill("solid", fgColor=RISK_RENG["Orta"])

    # Borc siyahısı
    if len(b):
        b = b.sort_values(["nov", "ad", "tarix"])
        ws = _vereq(wb, "Borc_siyahısı", ["Növ", "Kontragent", "VÖEN", "Qaimə", "Tarix", "Cəmi", "Ödənilib", "Qalıq",
                                          "Gün"], [10, 22, 13, 10, 11, 13, 13, 13, 7],
                    [["Debitor" if r.nov == "Satış" else "Kreditor", r.ad, r.voen, r.qaime, r.tarix, r.cem, r.odenilib,
                      r.qaliq, (dovr_sonu - r.tarix).days] for r in b.itertuples()],
                    reqem_sutunlari=(6, 7, 8), tarix_sutunlari=(5,))

        # Rəhbərlik hesabatı
        cem = b.groupby(["nov", "ad"], as_index=False)["qaliq"].sum().sort_values(["nov", "qaliq"], ascending=[False, False])
        _vereq(wb, "Rəhbərlik", ["Növ", "Kontragent", "Qalıq borc, AZN"], [12, 26, 16],
               [["Debitor" if r.nov == "Satış" else "Kreditor", r.ad, r.qaliq] for r in cem.itertuples()],
               reqem_sutunlari=(3,))

        # Məktub qaralamaları (yalnız debitorlar)
        email = dict(zip(kontragentler["voen"], kontragentler["email"]))
        setirler = []
        for (voen, adi), g in b[b["nov"] == "Satış"].groupby(["voen", "ad"]):
            maddeler = "\n".join(f"  - {r.qaime}, {r.tarix:%d.%m.%Y}: qaimə {az(r.cem)} AZN, ödənilib {az(r.odenilib)} AZN, "
                                 f"qalıq {az(r.qaliq)} AZN" for r in g.itertuples())
            metn = (f"Hörmətli {adi} mühasibatlığı,\n\n"
                    f"{dovr_sonu:%d.%m.%Y} tarixinə {C.SIRKET_ADI} qarşısında borcunuz {az(g['qaliq'].sum())} AZN təşkil edir:\n"
                    f"{maddeler}\n\n"
                    f"Məlumatlar Sizin uçotunuzla uyğun gəlmirsə, {C.CAVAB_MUDDETI_IS_GUNU} iş günü ərzində bizə bildirməyinizi xahiş edirik.\n\n"
                    f"Hörmətlə,\n{C.SIRKET_ADI} mühasibatlığı")
            setirler.append([adi, email.get(voen, ""), f"{C.SIRKET_ADI}: {dovr_sonu:%d.%m.%Y} tarixinə borc üzləşdirməsi", metn])
        _vereq(wb, "Məktub_qaralamaları", ["Kontragent", "E-poçt", "Mövzu", "Mətn"], [22, 24, 40, 90], setirler)

    wb.save(fayl)
