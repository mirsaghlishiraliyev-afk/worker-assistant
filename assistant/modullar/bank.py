"""Bank daxilolmaları modulu: bank çıxarışındakı daxilolmaları təsnif edir və 1C üçün Excel hazırlayır.

Proses (mühasibin cavablarından):
  - Kapital Bank çıxarışı (Excel) gəlir, bir neçə hesab ola bilər;
  - daxilolmalar: sığorta haqları (xüsusi Excel ilə 1C-yə yüklənir), subroqasiya, xəzinədarlıqdan
    əmək haqqı və ezamiyyə, ekspert namizədlərinin təlim haqları (bunlar 1C-yə əl ilə işlənir);
  - ən çox səhv: tarixin düzgün göstərilməməsi, pul vəsaitinin maddəsinin səhv seçilməsi.

Sistem yalnız TƏKLİF edir: hər sətir üçün kateqoriya, maddə, etibar dərəcəsi və səbəb. Tanımadığını
"Yoxlanmalı"ya salır. Mühasibin düzəlişi yadda saxlanır (data/bank_qaydalari.json): eyni ödəyici
növbəti dəfə avtomatik tanınır.
"""
import datetime as dt
import json
import pathlib
import re
from collections import Counter
from urllib.parse import quote

import pandas as pd

from uzlasma.config import az

from .excel import _reqem, norm

# ======================= ayarlar =======================
# Pul vəsaitinin maddələri 1C-dəki adlarla eyni olmalıdır: mühasib çatda dəyişə bilər
# («subroqasiya maddəsi "..." olsun»), dəyişiklik data/bank_qaydalari.json-da saxlanır.
KATEQORIYALAR = [
    {"kod": "sigorta_haqqi", "ad": "Sığorta haqları", "madde": "Sığorta haqları", "emal": "1C-yə Excel ilə",
     "acar": [r"sigort", r"sgort", r"sigorta haq", r"\bpolis", r"heyvan", r"\bekin", r"bitki", r"страхов", r"premium", r"\bsh\b"]},
    {"kod": "subroqasiya", "ad": "Subroqasiya", "madde": "Subroqasiya üzrə daxilolmalar", "emal": "1C-yə əl ilə",
     "acar": [r"subroq", r"reqres", r"regres", r"zererin evez", r"vurulmus zerer", r"ziyanin odenil", r"i\.i\.n\b", r"icra is",
              r"суброгац", r"регресс"]},
    {"kod": "emek_haqqi", "ad": "Əmək haqqı (xəzinədarlıq)", "madde": "Əmək haqqı üzrə daxilolmalar", "emal": "1C-yə əl ilə",
     "acar": [r"emek haqq", r"emekhaqq", r"\bmaas", r"\be/h\b", r"заработн"], "xezine": True},
    {"kod": "ezamiyye", "ad": "Ezamiyyə (xəzinədarlıq)", "madde": "Ezamiyyə üzrə daxilolmalar", "emal": "1C-yə əl ilə",
     "acar": [r"ezam", r"komandir"], "xezine": True},
    {"kod": "telim", "ad": "Təlim haqları", "madde": "Təlim haqları üzrə daxilolmalar", "emal": "1C-yə əl ilə",
     "acar": [r"t[ae]lim", r"tehsil haqq", r"\bkurs", r"namized", r"konsultasiya", r"qiymetlendirme", r"imtahan", r"sertifikat",
              r"обучен"]},
    {"kod": "qaytarilma", "ad": "Vəsaitin qaytarılması", "madde": "Qaytarılan vəsaitlər üzrə daxilolmalar", "emal": "1C-yə əl ilə",
     "acar": [r"qaytarilma", r"geri qaytar", r"artiq odenil"]},
    {"kod": "depozit", "ad": "Depozitin qaytarılması", "madde": "Depozit məbləğinin qaytarılması", "emal": "1C-yə əl ilə",
     "acar": [r"depozit mebleg", r"depozitin qaytar", r"depozit uzre esas"]},
    {"kod": "faiz", "ad": "Depozit faizləri", "madde": "Faiz gəlirləri", "emal": "1C-yə əl ilə",
     "acar": [r"faiz"]},
    # avtomatik təyin olunmur: yalnız mühasib seçir (sistemin sualına cavab olaraq)
    {"kod": "diger", "ad": "Digər daxilolmalar", "madde": "Digər daxilolmalar", "emal": "1C-yə əl ilə", "acar": []},
]
ZEIF = {"subroqasiya", "qaytarilma"}  # bu kateqoriyalar həmişə «Orta» etibarla təklif olunur (mühasib baxsın)
YOXLANMALI = {"kod": "yoxlanmali", "ad": "Yoxlanmalı", "madde": "", "emal": "mühasib təyin etməlidir"}
XEZINE = [r"xezinedar", r"\bdxa\b", r"казначейств", r"maliyye nazirliyi"]
XEZINE_HESAB = re.compile(r"ctre", re.I)  # Dövlət Xəzinədarlığının hesabları (AZ..CTRE..)
BANK_ADI = re.compile(r"kapital|\bbank\b", re.I)  # ödəyici yerində vasitəçi bank yazılır
OZ_ADI = [r"aqrar\s*sigorta\s*fondu", r"aqr\s*arsigorta\s*fondu"]  # Fondun öz adı təyinatda kateqoriya deyil
DOZUM_TEKRAR = 0.01

# 1C-yə yüklənən sığorta haqları faylının sütunları. Mühasibin şablonu gələndə bu siyahı ona uyğunlaşdırılacaq.
# Mühasibin göndərdiyi «Sığorta haqqı inteqrasiya faylı» şablonu (Sheet1, bu sütunlar bu sırada)
SUTUNLAR_1C = ["TARİX", "Sığortalı Fin/Voen", "Sığorta müqaviləsinin nömrəsi", "HESAB NÖMRƏSİ", "MƏBLƏĞ", "Pul vəsaitinin maddəsi"]
VEREQ_1C = "Sheet1"

QAYDA_FAYLI = None  # state.py təyin edir: data/bank_qaydalari.json
HESABAT_QOVLUGU = None  # state.py təyin edir: data/hesabatlar

# Çıxarışın sütun adları (bankdan banka fərqlənir). Sıra vacibdir: "Ödəyicinin VÖEN-i" əvvəl VÖEN kimi tanınsın.
ROLLAR = [
    ("eks", [r"eks hesab", r"qarsi hesab", r"kontragentin hesab", r"корреспондент.*сч"]),  # «IBAN / ad / VÖEN» bir xanada
    ("balans", [r"balans", r"balance", r"остаток", r"qaliq"]),
    ("voen", [r"voen", r"инн", r"\btin\b", r"vergi .*nomre"]),
    ("hesab", [r"iban", r"hesab", r"account", r"сч[её]т"]),
    ("valyuta_tarixi", [r"valyutala", r"value date", r"дата валют"]),
    ("tarix", [r"tarix", r"\bdate\b", r"дата"]),
    ("medaxil", [r"medaxil", r"kredit", r"daxilolma", r"daxil olan", r"credit", r"приход", r"кредит", r"\bcr\b"]),
    ("mexaric", [r"mexaric", r"debet", r"debit", r"расход", r"дебет", r"cixan", r"\bdr\b"]),
    ("mebleg", [r"mebleg", r"amount", r"сумма"]),
    ("teyinat", [r"teyinat", r"aciqlama", r"purpose", r"назначен", r"description", r"izah", r"qeyd", r"detal"]),
    ("odeyici", [r"odeyici", r"gonderen", r"kontragent", r"qarsi teref", r"payer", r"плательщ", r"корреспондент", r"\bad\b", r"adi$", r"musteri"]),
]

# Sığorta müqaviləsinin nömrəsi HS- və ya BS- ilə başlayır (məs. «HS-2026113347-QS»). Yalnız belə sətirlər 1C inteqrasiyasına düşür.
MUQAVILE_RE = re.compile(r"\b((?:HS|BS)-\d{4,}(?:-[A-Z]{1,4})?)\b", re.I)
IBAN_RE = re.compile(r"\bAZ\d{2}[A-Z]{4}[A-Z0-9]{20}\b")


# ======================= qaydalar (öyrənmə) =======================
def _qaydalar():
    if QAYDA_FAYLI and pathlib.Path(QAYDA_FAYLI).exists():
        try:
            d = json.loads(pathlib.Path(QAYDA_FAYLI).read_text(encoding="utf-8"))
            return {"odeyici": d.get("odeyici", {}), "madde": d.get("madde", {})}
        except (ValueError, OSError):
            pass
    return {"odeyici": {}, "madde": {}}


def _qaydalar_yaz(d):
    if QAYDA_FAYLI:
        pathlib.Path(QAYDA_FAYLI).parent.mkdir(parents=True, exist_ok=True)
        pathlib.Path(QAYDA_FAYLI).write_text(json.dumps(d, ensure_ascii=False, indent=1), encoding="utf-8")


def _kat(kod):
    return next((k for k in KATEQORIYALAR if k["kod"] == kod), YOXLANMALI)


def _madde(kod):
    return _qaydalar()["madde"].get(kod) or _kat(kod)["madde"]


def _odeyici_acari(voen, odeyici):
    v = re.sub(r"\D", "", str(voen or ""))
    return f"voen:{v}" if len(v) >= 7 else (f"ad:{norm(odeyici).strip()}" if str(odeyici or "").strip() else None)


# ======================= çıxarışı oxumaq =======================
def _rol(basliq):
    n = norm(basliq).strip()
    for rol, ifadeler in ROLLAR:
        if any(re.search(i, n) for i in ifadeler):
            return rol
    return None


def _basliq_tap(df):
    """İlk 40 sətirdə başlıq sətrini tapır: ən azı tarix, məbləğ (mədaxil) və təyinat olmalıdır."""
    for i in range(min(40, len(df))):
        rollar = {}
        for j, x in enumerate(df.iloc[i].tolist()):
            if isinstance(x, str) and x.strip():
                r = _rol(x)
                if r and r not in rollar:
                    rollar[r] = j
        if "tarix" in rollar and ("medaxil" in rollar or "mebleg" in rollar) and "teyinat" in rollar:
            return i, rollar
    return None, None


def _tarix(x):
    if x is None or (isinstance(x, float) and pd.isna(x)):
        return None
    if isinstance(x, (dt.datetime, pd.Timestamp)):
        return pd.Timestamp(x).to_pydatetime().date()
    if isinstance(x, dt.date):
        return x
    if isinstance(x, (int, float)) and 20000 < x < 80000:  # Excel seriya nömrəsi
        return (dt.datetime(1899, 12, 30) + dt.timedelta(days=float(x))).date()
    t = pd.to_datetime(str(x).strip(), dayfirst=True, errors="coerce", format="mixed")
    return None if pd.isna(t) else t.date()


MAKS_VEREQ = 10  # bank çıxarışında bir-neçə vərəq olur; 1C bazası kimi yüzlərlə vərəqli fayl çıxarış deyil


def taniyir(fayl):
    """Fayl bank çıxarışına bənzəyirsə, [(vərəq, başlıq sətri, rollar)] namizədləri qaytarır; deyilsə []."""
    from openpyxl import load_workbook
    try:
        wb = load_workbook(fayl, read_only=True, data_only=True)
    except Exception:
        return []
    try:
        adlar = wb.sheetnames
        if len(adlar) > MAKS_VEREQ:
            adlar = [a for a in adlar if re.search(r"cixaris|statement|выписк|bank", norm(a))][:MAKS_VEREQ]
        out = []
        for ad in adlar:
            rows = [list(r) for _, r in zip(range(60), wb[ad].iter_rows(values_only=True))]
            i, rollar = _basliq_tap(pd.DataFrame(rows, dtype=object)) if rows else (None, None)
            # mədaxil/məxaric sütunu, ya da işarəli tək «Məbləğ» + çıxarış əlaməti (başlıqda «çıxarış», «Balans», «Əks hesab»)
            if i is None:
                continue
            ust = norm(" ".join(str(x) for r in rows[:i] for x in r if x is not None))
            cixaris = re.search(r"cixaris|statement|выписк", ust) or "balans" in rollar or "eks" in rollar
            if "medaxil" in rollar or "mexaric" in rollar or ("mebleg" in rollar and cixaris):
                out.append((ad, i, rollar))
        return out
    finally:
        wb.close()


def oxu(fayl, ad=None):
    """Bank çıxarışını oxuyur və daxilolmaları təsnif edir. Bank çıxarışı deyilsə None."""
    for vereq, i, rollar in taniyir(fayl):
        b = _oxu_vereq(fayl, vereq, i, rollar, ad)
        if b["setirler"]:  # daxilolması olmayan vərəq çıxarış sayılmır
            tesnif_et(b)
            return b
    return None


def _meta(df, i):
    """Başlıqdan yuxarıdakı məlumat: hesab nömrəsi, müştərinin (Fondun) adı, dövr."""
    m = {}
    for k in range(i):
        r = [x for x in df.iloc[k].tolist() if not (x is None or (isinstance(x, float) and pd.isna(x)))]
        for a, b in zip(r, r[1:]):
            n = norm(a)
            if re.search(r"hesab nomresi|account number", n) and "hesab" not in m:
                m["hesab"] = str(b).strip()
            elif re.search(r"musterinin adi|client.s name", n):
                m["musteri"] = str(b).strip()
            elif re.search(r"dovr ucun|for period", n):
                m["dovr"] = str(b).strip()
    return m


def _teyinatdan(teyinat, eks_ad):
    """Təyinatdan ödəyici, FİN/VÖEN və qısa təyinat çıxarır (Kapital Bank formatları)."""
    t = str(teyinat or "")
    out = {"odeyici": "", "fin": "", "voen": "", "hesab": "", "qisa": t.strip()}
    # ödəyicinin hesabı və VÖEN-i təyinatda: «DTR_IBAN=AZ..», «TAX=1234567890», «VÖEN: ...»
    h = re.search(r"DTR_IBAN=(AZ\w+)", t) or IBAN_RE.search(t)
    out["hesab"] = h.group(1) if h and h.groups() else (h.group(0) if h else "")
    v = re.search(r"(?:TAX=|V[ÖO]EN[:\s]*)(\d{10})\b", t, re.I)
    out["voen"] = v.group(1) if v else ""
    m = re.search(r"Teyinat:\s*([^,]*)", t)
    if m:  # SV:..,Teyinat: X,Fin:..,ASA:<ad>,...
        out["qisa"] = re.sub(r"\s+", " ", m.group(1)).strip()
        a = re.search(r"ASA:([^,]+)", t)
        f = re.search(r"Fin:\s*([A-Za-z0-9]{5,8})", t)
        out["odeyici"] = a.group(1).strip() if a else ""
        out["fin"] = f.group(1).upper() if f else ""
        return out
    m = re.search(r"DTR_NAME=([^;]+)", t)
    if m:  # AniPay: ani ödəniş
        out["odeyici"] = m.group(1).strip()
        q = re.search(r"TAX=[^;]*;([^;]+)", t)
        out["qisa"] = re.sub(r"\s+", " ", q.group(1)).strip() if q else "Ani ödəniş (AniPay), təyinat yoxdur"
        return out
    m = re.match(r"^(.+?)\s+-\s+(.+?)(?:,\s*Million\S*)?$", t.strip())
    if m and BANK_ADI.search(eks_ad or ""):  # «AD SOYAD - telim odenisi,Million...»
        out["odeyici"], out["qisa"] = m.group(1).strip(), m.group(2).strip()
        return out
    m = re.search(r"^([^,]+),([^,]+),([^,]+),(?:EmanatMpay|KassamAz|Million)", t.strip())
    if m:  # «Təhsil haqqı,TELIM ODENISI,AD SOYAD,EmanatMpay...»
        ucuncu = m.group(3).strip()
        if re.fullmatch(r"[A-Z0-9]{7}", ucuncu):
            out["fin"] = ucuncu
        else:
            out["odeyici"] = ucuncu
        out["qisa"] = m.group(1).strip() + ("" if any(re.search(o, norm(m.group(2))) for o in OZ_ADI) else ", " + m.group(2).strip())
    return out


def _oxu_vereq(fayl, vereq, i, rollar, ad):
    df = pd.read_excel(fayl, sheet_name=vereq, header=None, dtype=object)
    meta = _meta(df, i)
    basliqlar = [str(x) if x is not None else "" for x in df.iloc[i].tolist()]
    setirler = []
    for k in range(i + 1, len(df)):
        r = df.iloc[k].tolist()
        def al(rol):
            x = r[rollar[rol]] if rol in rollar and rollar[rol] < len(r) else None
            return None if x is None or (isinstance(x, float) and pd.isna(x)) else x
        if "medaxil" in rollar:
            mebleg = _reqem(al("medaxil"))
        else:
            mebleg = _reqem(al("mebleg"))
        if mebleg is None or pd.isna(mebleg) or mebleg <= 0:
            continue  # məxaric və ya boş/cəm sətri
        teyinat = str(al("teyinat") or "").strip()
        if re.match(r"^(cemi|yekun|итого|total|qaliq|остаток)", norm(teyinat)) and not al("odeyici"):
            continue
        odeyici, voen, hesab = str(al("odeyici") or "").strip(), re.sub(r"\.0$", "", str(al("voen") or "").strip()), str(al("hesab") or "").strip()
        if "eks" in rollar:  # «AZ.. / KAPITAL BANK ASC / 1234567891»
            p = [x.strip() for x in str(al("eks") or "").split(" / ")]
            hesab = hesab or (p[0] if p else "")
            odeyici = odeyici or (p[1] if len(p) > 1 else "")
            voen = voen or (re.sub(r"\D", "", p[2]) if len(p) > 2 else "")
        ty = _teyinatdan(teyinat, odeyici)
        vasiteci = bool(BANK_ADI.search(odeyici))
        setirler.append({
            "setir": k + 1, "tarix_xam": al("tarix"), "tarix": _tarix(al("tarix")),
            "valyuta_tarixi": _tarix(al("valyuta_tarixi")), "mebleg": round(float(mebleg), 2),
            # ödəyici yerində bank yazılıbsa (vasitəçi), əsl ödəyici təyinatdan götürülür
            "odeyici": ty["odeyici"] or odeyici, "voen": ty["voen"] or ("" if vasiteci else voen), "fin": ty["fin"],
            "eks_ad": odeyici, "hesab": hesab, "teyinat": ty["qisa"], "teyinat_xam": teyinat,
            "odeyici_hesab": ty["hesab"], "t_fin_voen": ty["fin"] or ty["voen"],  # təyinatdan: 1C inteqrasiyası üçün
        })
    b = {"ad": ad or pathlib.Path(fayl).name, "vereq": vereq, "basliq_setri": i + 1, "meta": meta,
         "sutunlar": {rol: basliqlar[j] for rol, j in rollar.items()}, "setirler": setirler}
    return b


# ======================= təsnifat =======================
def _tesnif(r, qaydalar):
    """Bir daxilolma üçün (kateqoriya kodu, etibar, səbəb)."""
    acar = _odeyici_acari(r.get("voen") or r.get("fin"), r["odeyici"])
    if acar and acar in qaydalar["odeyici"]:
        kod = qaydalar["odeyici"][acar]
        return kod, "Yüksək", "mühasib bu ödəyicini əvvəl belə təsnif edib"
    mq = MUQAVILE_RE.search(r.get("teyinat_xam") or r["teyinat"])
    if mq:
        return "sigorta_haqqi", "Yüksək", f"təyinatda sığorta müqaviləsinin nömrəsi {mq.group(1).upper()}"
    t, o = norm(r.get("teyinat_xam") or r["teyinat"]), norm(r.get("eks_ad") or r["odeyici"])
    for oz in OZ_ADI + ([re.escape(norm(r["_musteri"]))] if r.get("_musteri") else []):
        t = re.sub(oz, " ", t)  # Fondun öz adı ("Aqrar Sığorta Fondu") sığorta haqqı əlaməti deyil
    xezine = any(re.search(x, o) or re.search(x, t) for x in XEZINE) or bool(XEZINE_HESAB.search(r.get("hesab") or ""))
    xal = {}
    sebeb = {}
    for k in KATEQORIYALAR:
        tutan = [a for a in k["acar"] if re.search(a, t)]
        if tutan:
            xal[k["kod"]] = 3 + (2 if k.get("xezine") and xezine else 0) - (2 if k.get("xezine") and not xezine else 0)
            sebeb[k["kod"]] = "təyinatda «" + re.search(tutan[0], t).group(0).strip() + "»" + (" və ödəyici xəzinədarlıqdır" if k.get("xezine") and xezine else "")
    if not xal:
        if xezine:
            return "yoxlanmali", "", "ödəyici xəzinədarlıqdır, amma təyinatdan əmək haqqı və ya ezamiyyə olduğu bilinmir"
        return "yoxlanmali", "", ("ani ödəniş, təyinat yazılmayıb" if "anipay" in t else "təyinatda tanış söz yoxdur")
    sirali = sorted(xal.items(), key=lambda x: -x[1])
    if len(sirali) > 1 and sirali[0][1] == sirali[1][1]:
        return "yoxlanmali", "", "həm «" + _kat(sirali[0][0])["ad"] + "», həm «" + _kat(sirali[1][0])["ad"] + "» ola bilər"
    kod, x = sirali[0]
    if x <= 1:
        return "yoxlanmali", "", sebeb[kod] + ", amma ödəyici xəzinədarlıq deyil"
    # tək kateqoriya tutubsa — yüksək etibar; bir neçəsi tutub, biri üstün gəlibsə — orta
    tek = sum(1 for _, v in sirali if v > 1) == 1
    return kod, ("Yüksək" if x >= 5 or (tek and kod not in ZEIF) else "Orta"), sebeb[kod]


def tesnif_et(b):
    q = _qaydalar()
    musteri = (b.get("meta") or {}).get("musteri", "")
    for r in b["setirler"]:
        if r.get("el_ile"):
            continue
        r["_musteri"] = musteri
        r["kateqoriya"], r["etibar"], r["sebeb"] = _tesnif(r, q)
        m = MUQAVILE_RE.search(r.get("teyinat_xam") or r["teyinat"])
        r["muqavile"] = m.group(1).upper() if m and r["kateqoriya"] == "sigorta_haqqi" else ""
    b["xeberdarliqlar"] = _yoxla(b)


def _yoxla(b):
    """Tarix və təkrar yoxlamaları (mühasibin dediyi ən çox səhv: tarixin düzgün göstərilməməsi)."""
    out = []
    ss = b["setirler"]
    aylar = Counter((r["tarix"].year, r["tarix"].month) for r in ss if r["tarix"])
    esas = aylar.most_common(1)[0][0] if aylar else None
    bu_gun = dt.date.today()
    for r in ss:
        if not r["tarix"]:
            out.append((r["setir"], "Tarix oxunmur", f"«{r['tarix_xam']}» tarix kimi tanınmadı — 1C-də tarixi əl ilə yoxlayın"))
        elif r["tarix"] > bu_gun:
            out.append((r["setir"], "Gələcək tarix", f"{r['tarix']:%d.%m.%Y} — çıxarışda səhv ola bilər"))
        elif esas and (r["tarix"].year, r["tarix"].month) != esas:
            out.append((r["setir"], "Başqa ay", f"{r['tarix']:%d.%m.%Y} çıxarışın əsas ayından ({esas[1]:02d}.{esas[0]}) fərqlidir — 1C-də hansı dövrə düşdüyünü yoxlayın"))
        if r["valyuta_tarixi"] and r["tarix"] and r["valyuta_tarixi"] != r["tarix"]:
            out.append((r["setir"], "Tarixlər fərqlidir", f"əməliyyat {r['tarix']:%d.%m.%Y}, valyutalaşma {r['valyuta_tarixi']:%d.%m.%Y} — 1C-də hansının götürüldüyünü dəqiqləşdirin"))
    gorulen = {}
    for r in ss:
        acar = (r["tarix"], round(r["mebleg"] / DOZUM_TEKRAR), norm(r["odeyici"]), norm(r["teyinat"]))
        if acar in gorulen:
            out.append((r["setir"], "Təkrar ola bilər", f"{gorulen[acar]}-ci sətirlə eyni (tarix, məbləğ, ödəyici, təyinat)"))
        else:
            gorulen[acar] = r["setir"]
    return sorted(out)


# ======================= alətlər =======================
def _yoxdur():
    return {"metn": "Bank çıxarışı yüklənməyib. 📎 ilə Kapital Bankın Excel çıxarışını əlavə edin.", "ui": None}


def _f(x):
    return az(x)


def bank_xulase(s):
    b = getattr(s, "bank", None)
    if not b:
        return _yoxdur()
    ss = b["setirler"]
    kartlar, hisseler = [], []
    for k in KATEQORIYALAR + [YOXLANMALI]:
        sec = [r for r in ss if r["kateqoriya"] == k["kod"]]
        if not sec:
            continue
        cem = sum(r["mebleg"] for r in sec)
        kartlar.append({"ad": k["ad"], "deyer": f"{_f(cem)} ₼", "alt": f"{len(sec)} daxilolma · {k['emal']}"})
        hisseler.append(f"{k['ad']}: {len(sec)} ({_f(cem)})")
    xb = len(b["xeberdarliqlar"])
    if xb:
        kartlar.append({"ad": "Xəbərdarlıq", "deyer": str(xb), "alt": "tarix, təkrar"})
    yox = sum(1 for r in ss if r["kateqoriya"] == "yoxlanmali")
    cem = sum(r["mebleg"] for r in ss)
    metn = (f"«{b['ad']}»: {len(ss)} daxilolma, cəmi {_f(cem)} AZN. " + "; ".join(hisseler) + "." +
            (f" {yox} daxilolmanı tanımadım — «Yoxlanmalıları göstər» yazın." if yox else "") +
            (f" Əmin olmadığım {len(_sual_setirleri(b))} daxilolma var — «Sualları ver» yazın, bir-bir soruşacağam." if _sual_setirleri(b) else "") +
            (f" {xb} xəbərdarlıq var (tarix, təkrar)." if xb else "") +
            " 1C üçün faylı almaq üçün yazın: «1C üçün Excel hazırla».")
    return {"metn": metn, "ui": {"tip": "kartlar", "kartlar": kartlar}}


def _kod_tap(soz):
    n = norm(soz or "")
    if not n:
        return None
    if re.search(r"yoxlan|taninma|bilinme|namelum|tanimadig", n):
        return "yoxlanmali"
    for k in KATEQORIYALAR:
        if n == k["kod"] or norm(k["ad"]).startswith(n[:6]) or any(re.search(a, n) for a in k["acar"]):
            return k["kod"]
    return None


def bank_siyahi(s, kateqoriya=None, n=50):
    b = getattr(s, "bank", None)
    if not b:
        return _yoxdur()
    if kateqoriya and norm(kateqoriya).startswith("xeberdar"):
        xs = b["xeberdarliqlar"]
        setirler = [[str(a), c, d] for a, c, d in xs[:n]]
        return {"metn": f"{len(xs)} xəbərdarlıq var." if xs else "Xəbərdarlıq yoxdur: tarixlər qaydasındadır, təkrar görünmür.",
                "ui": {"tip": "cedvel", "basliq": "Xəbərdarlıqlar", "sutunlar": ["Sətir", "Növ", "İzah"], "setirler": setirler} if xs else None}
    kod = _kod_tap(kateqoriya) if kateqoriya else None
    if kateqoriya and not kod:
        return {"metn": f"«{kateqoriya}» kateqoriyası yoxdur. Kateqoriyalar: " + ", ".join(k["ad"] for k in KATEQORIYALAR) + ", Yoxlanmalı.", "ui": None}
    sec = [r for r in b["setirler"] if not kod or r["kateqoriya"] == kod]
    k = _kat(kod) if kod else None
    sutunlar = ["Sətir", "Tarix", "Ödəyici", "Məbləğ", "Təyinat", "Kateqoriya", "Maddə", "Etibar", "Səbəb"]  # noqa
    setirler = [[str(r["setir"]), f"{r['tarix']:%d.%m.%Y}" if r["tarix"] else str(r["tarix_xam"] or ""), r["odeyici"], _f(r["mebleg"]),
                 r["teyinat"][:90], _kat(r["kateqoriya"])["ad"], _madde(r["kateqoriya"]), r["etibar"] or "—", r["sebeb"]] for r in sec[:n]]
    if kod == "yoxlanmali":  # boş sütunlar (kateqoriya, maddə, etibar) göstərilmir
        sutunlar = ["Sətir", "Tarix", "Ödəyici", "Məbləğ", "Təyinat", "Niyə tanınmadı"]
        setirler = [[x[0], x[1], x[2], x[3], x[4], x[8]] for x in setirler]
    cem = sum(r["mebleg"] for r in sec)
    ad = k["ad"] if k else "Bütün daxilolmalar"
    metn = f"{ad}: {len(sec)} daxilolma, cəmi {_f(cem)} AZN" + (f", ilk {n}-i göstərilir" if len(sec) > n else "") + "."
    if kod == "yoxlanmali" and sec:
        metn += " Düzəltmək üçün yazın, məs.: «" + str(sec[0]["setir"]) + "-ci sətir subroqasiyadır». Ödəyici yadda saxlanacaq."
    return {"metn": metn, "ui": {"tip": "cedvel", "basliq": f"{ad} · {_f(cem)} ₼", "sutunlar": sutunlar, "setirler": setirler}}


def bank_teyin(s, setir, kateqoriya, yadda_saxla=True):
    """Mühasibin düzəlişi: sətrin kateqoriyasını dəyişir və ödəyicini yadda saxlayır."""
    b = getattr(s, "bank", None)
    if not b:
        return _yoxdur()
    kod = _kod_tap(kateqoriya)
    if not kod or kod == "yoxlanmali":
        return {"metn": f"«{kateqoriya}» tanınmadı. Kateqoriyalar: " + ", ".join(k["ad"] for k in KATEQORIYALAR) + ".", "ui": None}
    r = next((x for x in b["setirler"] if x["setir"] == int(setir)), None)
    if not r:
        return {"metn": f"{setir}-ci sətir daxilolmalar arasında yoxdur (sətir nömrəsi Excel-dəki kimidir).", "ui": None}
    acar = _odeyici_acari(r["voen"], r["odeyici"])
    elave = ""
    if yadda_saxla and acar:
        q = _qaydalar()
        q["odeyici"][acar] = kod
        _qaydalar_yaz(q)
        evvel = {x["setir"]: x["kateqoriya"] for x in b["setirler"]}
        tesnif_et(b)
        diger = [x for x in b["setirler"] if x["setir"] != r["setir"] and evvel[x["setir"]] != x["kateqoriya"]]
        if diger:
            elave = f" Eyni ödəyicinin daha {len(diger)} daxilolması da «{_kat(kod)['ad']}» kateqoriyasına keçdi."
        elave += " Bu ödəyici növbəti çıxarışlarda avtomatik tanınacaq."
    else:
        r.update(kateqoriya=kod, etibar="Yüksək", sebeb="mühasib təyin edib", el_ile=True)
    r["cavablandi"] = True
    metn = (f"Oldu: {r['setir']}-ci sətir ({r['odeyici'] or 'ödəyici yoxdur'}, {_f(r['mebleg'])}) — «{_kat(kod)['ad']}», "
            f"maddə: «{_madde(kod)}».{elave}")
    if b.get("aktiv_sual") is not None:  # sual-cavab rejimindəyik: növbəti sualı ver
        novbeti = bank_sual(s)
        return {"metn": metn + "\n\n" + novbeti["metn"], "ui": novbeti["ui"]}
    return {"metn": metn, "ui": None}


# ======================= mühasibə suallar =======================
def _sual_setirleri(b):
    """Sistemin əmin olmadığı daxilolmalar: tanınmayanlar və «Orta» etibarlılar (cavablandırılmamış, keçilməmiş)."""
    return [r for r in b["setirler"] if (r["kateqoriya"] == "yoxlanmali" or r["etibar"] == "Orta")
            and not r.get("cavablandi") and not r.get("kecildi")]


def bank_sual(s):
    """Növbəti sualı verir: sətir haqqında məlumat, sistemin təxmini və cavab variantları (düymələr)."""
    b = getattr(s, "bank", None)
    if not b:
        return _yoxdur()
    ss = _sual_setirleri(b)
    if not ss:
        b["aktiv_sual"] = None
        return {"metn": "Sualım qalmadı: bütün daxilolmalar təsnif olunub. 1C faylı üçün yazın: «1C üçün Excel hazırla».", "ui": None}
    r = ss[0]
    b["aktiv_sual"] = r["setir"]
    cem = len(_sual_setirleri(b)) + sum(1 for x in b["setirler"] if x.get("cavablandi") or x.get("kecildi"))
    no = cem - len(ss) + 1
    tarix = f"{r['tarix']:%d.%m.%Y}" if r["tarix"] else str(r["tarix_xam"] or "")
    if r["kateqoriya"] == "yoxlanmali":
        texmin = f"Tanıya bilmədim: {r['sebeb']}."
    else:
        texmin = f"Məncə bu «{_kat(r['kateqoriya'])['ad']}»dir ({r['sebeb']}), amma əmin deyiləm."
    metn = (f"Sual {no}/{cem}: {r['setir']}-ci sətir — {tarix}, **{_f(r['mebleg'])} AZN**, ödəyici: {r['odeyici'] or 'yazılmayıb'}.\n"
            f"Təyinat: «{r['teyinat'][:160]}».\n{texmin} Bu daxilolma nədir? Düyməni seçin və ya öz sözlərinizlə yazın.")
    variantlar = []
    if r["kateqoriya"] != "yoxlanmali":
        variantlar.append({"etiket": f"✓ Bəli, {_kat(r['kateqoriya'])['ad'].lower()}", "mesaj": f"{r['setir']}-ci sətir — {_kat(r['kateqoriya'])['ad']} olsun",
                           "esas": True})
    for k in KATEQORIYALAR:
        if k["kod"] != r["kateqoriya"]:
            variantlar.append({"etiket": k["ad"], "mesaj": f"{r['setir']}-ci sətir — {k['ad']} olsun"})
    variantlar.append({"etiket": "Sonra baxaram", "mesaj": f"{r['setir']}-ci sətri keç", "ikinci": True})
    return {"metn": metn, "ui": {"tip": "sual", "setir": r["setir"], "variantlar": variantlar}}


def bank_kec(s, setir=None):
    b = getattr(s, "bank", None)
    if not b:
        return _yoxdur()
    setir = setir if setir is not None else b.get("aktiv_sual")
    r = next((x for x in b["setirler"] if x["setir"] == int(setir)), None) if setir is not None else None
    if r:
        r["kecildi"] = True
    novbeti = bank_sual(s)
    return {"metn": (f"{setir}-ci sətri sonraya saxladım (Yoxlanmalı/Orta olaraq qalır).\n\n" if r else "") + novbeti["metn"],
            "ui": novbeti["ui"]}


def bank_madde(s, kateqoriya, madde):
    """Kateqoriyanın pul vəsaitinin maddəsini 1C-dəki adla eyniləşdirir."""
    kod = _kod_tap(kateqoriya)
    if not kod or kod == "yoxlanmali":
        return {"metn": f"«{kateqoriya}» kateqoriyası tanınmadı.", "ui": None}
    q = _qaydalar()
    q["madde"][kod] = madde.strip()
    _qaydalar_yaz(q)
    return {"metn": f"Oldu: «{_kat(kod)['ad']}» üçün maddə bundan sonra «{madde.strip()}» yazılacaq.", "ui": None}


# ======================= inteqrasiya şablonu =======================
def sablon_mi(fayl):
    """Fayl «Sığorta haqqı inteqrasiya faylı» şablonudursa: (vərəq, başlıq sətri, {sütun adı: sütun №}); deyilsə None."""
    from openpyxl import load_workbook
    lazim = {norm(c): c for c in SUTUNLAR_1C}
    try:
        wb = load_workbook(fayl, read_only=True)
    except Exception:
        return None
    try:
        if len(wb.sheetnames) > MAKS_VEREQ:
            return None
        for ws in wb.worksheets:
            for i, row in zip(range(1, 11), ws.iter_rows(max_row=10, values_only=True)):
                cols = {lazim[norm(x).strip()]: j for j, x in enumerate(row, 1) if isinstance(x, str) and norm(x).strip() in lazim}
                if len(cols) >= 4:
                    return ws.title, i, cols
    finally:
        wb.close()
    return None


def _sablonu_doldur(sablon, setirler, qovluq, vaxt):
    """Mühasibin öz şablon faylını açır, başlığa görə sütunları tapır, mövcud sətirlərin altına yazır."""
    from openpyxl import load_workbook
    vereq, bas, cols = sablon_mi(sablon["yol"])
    wb = load_workbook(sablon["yol"])
    ws = wb[vereq]
    son = bas
    for r in range(ws.max_row, bas, -1):  # mövcud doldurulmuş sətirlər qalır
        if any(ws.cell(r, c).value not in (None, "") for c in cols.values()):
            son = r
            break
    for i, deyerler in enumerate(setirler, son + 1):
        for ad, j in cols.items():
            c = ws.cell(i, j, deyerler[SUTUNLAR_1C.index(ad)])
            if ad == "TARİX":
                c.number_format = "DD.MM.YYYY"
            elif ad == "MƏBLƏĞ":
                c.number_format = "0.00"
    kok = pathlib.Path(sablon["ad"]).stem
    ad = f"{kok} (doldurulub {vaxt[:8]}_{vaxt[9:]}).xlsx"
    wb.save(qovluq / ad)
    return ad, son - bas


def bank_hesabat(s, muqavilesiz=False):
    """1C üçün Excel: sığorta haqları (yükləmə faylı), əl ilə işlənənlər, yoxlanmalılar, xəbərdarlıqlar."""
    b = getattr(s, "bank", None)
    if not b:
        return _yoxdur()
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    xu = wb.active
    xu.title = "Xülasə"
    xu.append(["Fayl", b["ad"]])
    xu.append(["Hazırlanıb", f"{dt.datetime.now():%d.%m.%Y %H:%M}"])
    xu.append(["Qeyd", "Sistem yalnız təklif edir. 1C-yə yükləməzdən əvvəl yoxlayın."])
    xu.append([])
    xu.append(["Kateqoriya", "Daxilolma", "Məbləğ", "1C-yə", "Pul vəsaitinin maddəsi"])
    for k in KATEQORIYALAR + [YOXLANMALI]:
        sec = [r for r in b["setirler"] if r["kateqoriya"] == k["kod"]]
        xu.append([k["ad"], len(sec), round(sum(r["mebleg"] for r in sec), 2), k["emal"], _madde(k["kod"])])
    xu.append(["Cəmi", len(b["setirler"]), round(sum(r["mebleg"] for r in b["setirler"]), 2)])

    # təkrar ola biləcək sətirlər 1C-yə yüklənmir: «Yoxlanmalı»ya düşür
    tekrar = {x[0]: x[2] for x in b["xeberdarliqlar"] if x[1] == "Təkrar ola bilər"}

    def vereq(ad, kod, sutunlar, setir):
        ws = wb.create_sheet(ad[:31])
        ws.append(sutunlar)
        for r in b["setirler"]:
            if kod == "yoxlanmali":
                if r["kateqoriya"] == kod or r["setir"] in tekrar:
                    ws.append(setir(r))
            elif r["kateqoriya"] == kod and r["setir"] not in tekrar:
                ws.append(setir(r))
        return ws

    def inteqrasiya(r):  # 1C inteqrasiya faylına müqavilə nömrəsi (HS-/BS-) olan sığorta haqları düşür;
        # mühasib istəsə (muqavilesiz=True), müqaviləsizlər də — nömrə sütunu boş qalır
        return r["kateqoriya"] == "sigorta_haqqi" and (r["muqavile"] or muqavilesiz) and r["setir"] not in tekrar

    def setir_1c(r):  # FİN/VÖEN və ödəyicinin hesabı təyinatdan götürülür
        d = {"TARİX": r["tarix"], "Sığortalı Fin/Voen": r.get("t_fin_voen") or r.get("fin") or r["voen"],
             "Sığorta müqaviləsinin nömrəsi": r["muqavile"], "HESAB NÖMRƏSİ": r.get("odeyici_hesab", ""),
             "MƏBLƏĞ": r["mebleg"], "Pul vəsaitinin maddəsi": _madde(r["kateqoriya"])}
        return [d.get(c, "") for c in SUTUNLAR_1C]

    def qeyd(r):
        return "; ".join(x for x in [("FİN/VÖEN təyinatda yoxdur" if not (r.get("t_fin_voen") or r.get("fin") or r["voen"]) else ""),
                                     ("ödəyicinin hesabı təyinatda yoxdur" if not r.get("odeyici_hesab") else "")] if x)
    # yoxlama faylında: 1C sütunları + mühasibin yoxlaması üçün ödəyici və təyinat
    w1 = wb.create_sheet("1C - Sığorta haqları")
    w1.append(SUTUNLAR_1C + ["Çıxarış sətri", "Ödəyici", "Təyinat", "Qeyd"])
    for r in b["setirler"]:
        if inteqrasiya(r):
            w1.append(setir_1c(r) + [r["setir"], r["odeyici"], r["teyinat"], qeyd(r)])
    for c in w1[1][len(SUTUNLAR_1C):]:
        c.fill = PatternFill("solid", fgColor="E8E8E8")
    w2 = wb.create_sheet("Sığorta haqları (müqaviləsiz)")
    w2.append(["Çıxarış sətri", "Tarix", "Ödəyici", "FİN/VÖEN", "Məbləğ", "Təyinat", "Qeyd"])
    for r in b["setirler"]:
        if r["kateqoriya"] == "sigorta_haqqi" and not r["muqavile"] and r["setir"] not in tekrar:
            if muqavilesiz:
                w2.append([r["setir"], r["tarix"], r["odeyici"], r.get("fin") or r["voen"], r["mebleg"], r["teyinat"],
                           "müqavilə nömrəsi yoxdur — mühasibin istəyi ilə inteqrasiyaya salınıb (nömrə sütunu boş)"])
                continue
            w2.append([r["setir"], r["tarix"], r["odeyici"], r.get("fin") or r["voen"], r["mebleg"], r["teyinat"],
                       "müqavilə nömrəsi (HS-/BS-) yoxdur — inteqrasiyaya düşmür, 1C-yə əl ilə"])
    el = ["Çıxarış sətri", "Tarix", "Ödəyici", "FİN/VÖEN", "Məbləğ", "Təyinat", "Pul vəsaitinin maddəsi", "Etibar", "Səbəb"]
    for k in KATEQORIYALAR[1:]:
        vereq(k["ad"], k["kod"], el, lambda r: [r["setir"], r["tarix"], r["odeyici"], r.get("fin") or r["voen"], r["mebleg"], r["teyinat"],
                                                 _madde(r["kateqoriya"]), r["etibar"], r["sebeb"]])
    ws = vereq("Yoxlanmalı", "yoxlanmali", ["Çıxarış sətri", "Tarix", "Ödəyici", "FİN/VÖEN", "Məbləğ", "Təyinat", "Səbəb", "Kateqoriya (doldurun)"],
               lambda r: [r["setir"], r["tarix"], r["odeyici"], r.get("fin") or r["voen"], r["mebleg"], r["teyinat"],
                          ("təkrar ola bilər: " + tekrar[r["setir"]]) if r["setir"] in tekrar else r["sebeb"],
                          "" if r["kateqoriya"] == "yoxlanmali" else _kat(r["kateqoriya"])["ad"] + "?"])
    for c in ws[1]:
        c.fill = PatternFill("solid", fgColor="FFE5B4")
    xw = wb.create_sheet("Xəbərdarlıqlar")
    xw.append(["Çıxarış sətri", "Növ", "İzah"])
    for x in b["xeberdarliqlar"]:
        xw.append(list(x))
    for w in wb.worksheets:
        for c in w[1 if w.title != "Xülasə" else 5]:
            c.font = Font(bold=True)
        for col in w.columns:
            w.column_dimensions[col[0].column_letter].width = min(60, max(10, max(len(str(c.value or "")) for c in col) + 2))
        for row in w.iter_rows(min_row=2):
            for c in row:
                if isinstance(c.value, dt.date):
                    c.number_format = "DD.MM.YYYY"
                elif isinstance(c.value, float):
                    c.number_format = "#,##0.00"
    qovluq = pathlib.Path(HESABAT_QOVLUGU or ".")
    qovluq.mkdir(parents=True, exist_ok=True)
    vaxt = f"{dt.datetime.now():%Y%m%d_%H%M%S}"
    ad = f"bank_daxilolmalari_{vaxt}.xlsx"
    wb.save(qovluq / ad)

    # 1C inteqrasiya faylı: mühasibin şablonu ilə eyni (Sheet1, yalnız şablon sütunları)
    wi = Workbook()
    si = wi.active
    si.title = VEREQ_1C
    si.append(SUTUNLAR_1C)
    sh = 0
    for r in b["setirler"]:
        if inteqrasiya(r):
            si.append(setir_1c(r))
            sh += 1
    for row in si.iter_rows(min_row=2):
        row[0].number_format = "DD.MM.YYYY"
        row[SUTUNLAR_1C.index("MƏBLƏĞ")].number_format = "0.00"
    for col in si.columns:
        si.column_dimensions[col[0].column_letter].width = max(12, max(len(str(c.value or "")) for c in col) + 2)
    ad_1c = f"sigorta_haqqi_inteqrasiya_{vaxt}.xlsx"
    sablon = getattr(s, "sablon", None)
    evvelki = 0
    if sablon:  # mühasibin yüklədiyi öz şablonu doldurulur (sütun sırası və formatı onun faylındakı kimi qalır)
        ad_1c, evvelki = _sablonu_doldur(sablon, [setir_1c(r) for r in b["setirler"] if inteqrasiya(r)], qovluq, vaxt)
    else:
        wi.save(qovluq / ad_1c)

    yox = sum(1 for r in b["setirler"] if r["kateqoriya"] == "yoxlanmali" or r["setir"] in tekrar)
    natamam = sum(1 for r in b["setirler"] if inteqrasiya(r) and qeyd(r))
    musz = sum(1 for r in b["setirler"] if r["kateqoriya"] == "sigorta_haqqi" and r["setir"] not in tekrar and not r["muqavile"])
    sablon_metn = (f" Yüklədiyiniz «{sablon['ad']}» şablonu dolduruldu" + (f" (əvvəlki {evvelki} sətir saxlanıb, yenilər altına yazılıb)" if evvelki else "") + "."
                   if sablon else "")
    if muqavilesiz:
        bas = (f"İki fayl hazırdır: «Sığorta haqqı inteqrasiya faylı» — {sh} sətir, o cümlədən müqavilə nömrəsi olmayan {musz} sətir "
               f"(nömrə sütunu boş); və yoxlama faylı (bütün kateqoriyalar" + (f", {yox} sətir «Yoxlanmalı»" if yox else "") + ").")
    else:
        bas = (f"İki fayl hazırdır: «Sığorta haqqı inteqrasiya faylı» — müqavilə nömrəsi (HS-/BS-) olan {sh} sətir, şablonunuzla eyni "
               f"sütunlar; və yoxlama faylı (bütün kateqoriyalar" + (f", {yox} sətir «Yoxlanmalı»" if yox else "") + ")." +
               (f" Müqavilə nömrəsi olmayan {musz} sığorta haqqı inteqrasiyaya düşmür — yoxlama faylında ayrıca vərəqdədir. "
                f"Onları da şablona salmaq üçün yazın: «Müqaviləsizləri də şablona sal»." if musz else ""))
    if sablon:  # əsas nəticə şablondur: əvvəl onu de
        bas = (f"«{sablon['ad']}» şablonunuz dolduruldu: {sh} yeni sətir" + (f" (əvvəlki {evvelki} sətir saxlanıb, yenilər altına yazılıb)" if evvelki else "") +
               ". Yanında yoxlama faylı da var." + bas[bas.index(")") + 2:] if ")" in bas else bas)
        sablon_metn = ""
    return {"metn": bas + sablon_metn +
                    (f" {natamam} sətirdə FİN/VÖEN və ya ödəyicinin hesabı təyinatda tapılmadı — «Qeyd» sütununa baxın." if natamam else ""),
            "ui": {"tip": "fayllar", "fayllar": [{"ad": ad_1c, "url": "/api/fayl/" + quote(ad_1c)}, {"ad": ad, "url": "/api/fayl/" + quote(ad)}]}}


# ======================= lokal (AI-sız) rejim =======================
def qayda(s, mesaj):
    """Bank sualını alətə çevirir: (alət, args) və ya None."""
    if not getattr(s, "bank", None):
        return None
    n = norm(mesaj)
    # "12-ci sətir subroqasiyadır", "sətir 12 sığorta haqqı"
    m = re.search(r"(\d+)\s*-?\s*(?:ci|cu|cı|cü|ci|nci|ncu|inci|uncu)?\s*set(?:ir|ri)|set(?:ir|ri)\s*(\d+)", n)
    if m and re.search(r"\bkec|sonra|atla", n):
        return "bank_kec", {"setir": int(m.group(1) or m.group(2))}
    if m and re.search(r"dir|dur|dır|dür|olsun|kimi|yaz|deyis|kecir|\bdi\b|aiddir", n):
        kod = None
        for k in KATEQORIYALAR:
            if any(re.search(a, n) for a in k["acar"]) or norm(k["ad"]).split()[0][:6] in n:
                kod = k["kod"]
                break
        if kod:
            return "bank_teyin", {"setir": int(m.group(1) or m.group(2)), "kateqoriya": kod}
    # aktiv sual varsa: «bəli», «xeyr», «bilmirəm» və ya sərbəst cavab («icra qaydasında subroqasiyadır»)
    aktiv = s.bank.get("aktiv_sual")
    if re.search(r"suallar|sual ver|soruş|sorus|cavablandir", n) and not m:
        return "bank_sual", {}
    if aktiv is not None and not m and len(n.split()) <= 25:
        r = next((x for x in s.bank["setirler"] if x["setir"] == aktiv), None)
        if re.search(r"^\s*(beli|he|hə|duzdur|dogrudur|ok|razi|tesdiq)\b", n) and r and r["kateqoriya"] != "yoxlanmali":
            return "bank_teyin", {"setir": aktiv, "kateqoriya": r["kateqoriya"]}
        if re.search(r"bilmirem|sonra|kec\b|atla|emin deyil", n):
            return "bank_kec", {"setir": aktiv}
        for k in KATEQORIYALAR:
            if any(re.search(a, n) for a in k["acar"]) or norm(k["ad"]).split()[0][:6] in n:
                return "bank_teyin", {"setir": aktiv, "kateqoriya": k["kod"]}
        if re.search(r"^\s*(xeyr|yox|deyil)\b", n):
            return "bank_sual", {}
    # "subroqasiya maddəsi «...» olsun"
    sitat = re.search(r"[«\"']([^»\"']+)[»\"']", mesaj)
    if "madde" in n and sitat:
        kod = next((k["kod"] for k in KATEQORIYALAR if any(re.search(a, n.replace(norm(sitat.group(1)), "")) for a in k["acar"])), None)
        if kod:
            return "bank_madde", {"kateqoriya": kod, "madde": sitat.group(1)}
    if re.search(r"muqavilesiz", n):  # «Müqaviləsizləri də şablona sal» / «müqaviləsizləri salma»
        return "bank_hesabat", {"muqavilesiz": not re.search(r"salma|cixar|daxil etme|olmasin", n)}
    if re.search(r"\b1c\b|hesabat|excel|fayl hazirla|yukleme fayli|sablon|inteqrasiya|doldur", n):
        return "bank_hesabat", {}
    if re.search(r"xeberdarliq|tarix|tekrar", n):
        return "bank_siyahi", {"kateqoriya": "xeberdarliq"}
    if re.search(r"yoxlanmali|taninmayan|tanimadig|bilinmeyen|namelum", n):
        return "bank_siyahi", {"kateqoriya": "yoxlanmali"}
    for k in KATEQORIYALAR:
        if any(re.search(a, n) for a in k["acar"]):
            return "bank_siyahi", {"kateqoriya": k["kod"]}
    if re.search(r"hamisi|butun daxilolma|siyahi", n):
        return "bank_siyahi", {}
    if re.search(r"bank|daxilolma|cixaris|xulase|veziyyet|umumi|nece", n):
        return "bank_xulase", {}
    return None


# ======================= sxemlər =======================
_kateq = {"type": "string", "description": "sigorta_haqqi, subroqasiya, emek_haqqi, ezamiyye, telim, qaytarilma, depozit, faiz, diger, yoxlanmali və ya xeberdarliq"}
ALETLER = {
    "bank_xulase": (bank_xulase, "Bank çıxarışındakı daxilolmaların kateqoriyalar üzrə xülasəsi.", {}),
    "bank_siyahi": (bank_siyahi, "Kateqoriya üzrə daxilolmalar (və ya tanınmayanlar, xəbərdarlıqlar).",
                    {"kateqoriya": _kateq, "n": {"type": "integer"}}),
    "bank_teyin": (bank_teyin, "Mühasibin düzəlişi: sətrin kateqoriyasını dəyişir, ödəyicini yadda saxlayır.",
                   {"setir": {"type": "integer"}, "kateqoriya": _kateq, "yadda_saxla": {"type": "boolean"}}),
    "bank_sual": (bank_sual, "Sistemin əmin olmadığı növbəti daxilolma haqqında mühasibə sual verir (variant düymələri ilə).", {}),
    "bank_kec": (bank_kec, "Sualı sonraya saxlayır və növbəti sualı verir.", {"setir": {"type": "integer"}}),
    "bank_madde": (bank_madde, "Kateqoriyanın pul vəsaitinin maddəsini 1C-dəki adla eyniləşdirir.",
                   {"kateqoriya": _kateq, "madde": {"type": "string"}}),
    "bank_hesabat": (bank_hesabat, "1C üçün Excel: sığorta haqqı inteqrasiya faylı (HS-/BS- müqavilə nömrəli sətirlər) və yoxlama faylı. "
                     "muqavilesiz=true olsa, müqavilə nömrəsi olmayan sığorta haqları da inteqrasiyaya salınır.",
                     {"muqavilesiz": {"type": "boolean"}}),
}
MECBURI = {"bank_teyin": ["setir", "kateqoriya"], "bank_madde": ["kateqoriya", "madde"]}

MODUL = {
    "kod": "bank",
    "ad": "Bank daxilolmaları",
    "tesvir": "Bank çıxarışındakı daxilolmaları təsnif edir (sığorta haqqı, subroqasiya, əmək haqqı, ezamiyyə, təlim), "
              "maddəni təklif edir və 1C üçün Excel hazırlayır",
    "status": "aktiv",
    "teklifler": [
        {"s": "Bank daxilolmalarının xülasəsi", "b": "Daxilolmalar", "a": "kateqoriyalar üzrə"},
        {"s": "Yoxlanmalıları göstər", "b": "Tanınmayanlar", "a": "mühasib təyin etməlidir"},
        {"s": "1C üçün Excel hazırla", "b": "1C faylı", "a": "sığorta haqları + digərləri"},
    ],
}
